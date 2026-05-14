"""Verify Word-math parameter symbols stay attached to their explanations.

Run:
    uv run python scripts/test_word_math_symbol_legend.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

_ROOT = Path(__file__).resolve().parent.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from docx import Document  # noqa: E402

from scripts.make_word_math_equations_docx import main as make_sample  # noqa: E402
from translate.agent.docx_utils import extract_all_text, extract_math_texts, has_math  # noqa: E402
from translate.agent.graph import build_graph  # noqa: E402
from translate.agent.state import TranslationState  # noqa: E402


class WordMathStub:
    def complete(self, messages):
        user = messages[-1]["content"]
        if "Plan the English preamble" in user:
            return json.dumps({
                "korean_subject_span": "제어 장치",
                "english_noun_phrase": "control device",
                "independent_preamble": "A control device comprising:",
                "actor_phrase": "",
                "confidence": "high",
            })
        if "Translate Korean claim" in user:
            return json.dumps({
                "text": (
                    "A control device comprising:\n"
                    "an output unit satisfying the following mathematical expressions:\n"
                    "[EQUATION_1]\n"
                    "[EQUATION_2]\n"
                    "wherein A₁ + B₂/2 is a first combined input value; "
                    "B₂ − C is a second combined input value; "
                    "and C² + 1/2 is a correction value."
                ),
                "key_terms": [],
            })
        if "제어부는 다음 수학식들을 이용하여 출력값을 산출한다" in user:
            return json.dumps({
                "text": (
                    "The controller calculates an output value using the following "
                    "mathematical expressions:"
                ),
                "key_terms": [],
            })
        body_clauses = {
            "제1 복합 가중치": "α² + 1/2 is a first composite weight",
            "제2 복합 가중치": "β + γ/2 is a second composite weight",
            "보정 함수": "γ² − α is a correction function",
            "제1 입력값": "X₁ is a first input value",
            "제2 입력값": "Y₂ is a second input value",
            "제3 입력값": "Z₃ is a third input value",
        }
        for korean, text in body_clauses.items():
            if korean in user:
                return json.dumps({"text": text, "key_terms": []})
        if "[EQUATION_1]" in user and "α² + 1/2는" in user:
            return json.dumps({
                "text": (
                    "The controller calculates an output value using the following mathematical expressions:\n"
                    "[EQUATION_1]\n"
                    "[EQUATION_2]\n"
                    "[EQUATION_3]\n"
                    "where α² + 1/2 is a first composite weight; "
                    "β + γ/2 is a second composite weight; "
                    "γ² − α is a correction function; "
                    "X₁ is a first input value; Y₂ is a second input value; "
                    "and Z₃ is a third input value."
                ),
                "key_terms": [],
            })
        return json.dumps({"text": "Translated text.", "key_terms": []})


def main() -> None:
    make_sample()
    src = Path("data/sample_word_math_equations.docx")
    dst = Path("output/sample_word_math_equations_stub_en.docx")
    if dst.exists():
        dst.unlink()

    doc = Document(src)
    assert "α² + 1/2는 제1 복합 가중치" in extract_all_text(doc.paragraphs[6])
    assert "[EQUATION]는 제1 복합 가중치" not in extract_all_text(doc.paragraphs[6])
    assert "A₁ + B₂/2는 제1 조합 입력값" in extract_all_text(doc.paragraphs[12])
    assert "[EQUATION]는 제1 조합 입력값" not in extract_all_text(doc.paragraphs[12])

    state: TranslationState = {
        "input_path": src,
        "output_path": dst,
        "font": "Times New Roman",
        "review": False,
        "verbose": False,
        "delay": 0.0,
        "client": WordMathStub(),
    }
    build_graph().invoke(state)

    out = Document(dst)
    texts = [(p.text or "").strip() for p in out.paragraphs]
    all_text = "\n".join(extract_all_text(p).strip() for p in out.paragraphs)
    assert "A₁ + B₂/2 is a first combined input value" in all_text
    assert "B₂ − C is a second combined input value" in all_text
    assert "C² + 1/2 is a correction value" in all_text
    assert "α² + 1/2 is a first composite weight" in all_text
    assert "β + γ/2 is a second composite weight" in all_text
    assert "γ² − α is a correction function" in all_text

    body_sequence = []
    for p in out.paragraphs[:9]:
        text = extract_all_text(p).strip()
        if has_math(p) and text == "[EQUATION]":
            body_sequence.append("EQ")
        elif text.startswith((", where α²", "wherein α²", "where α²")):
            body_sequence.append("ALPHA")
    assert body_sequence[:4] == ["EQ", "EQ", "EQ", "ALPHA"], body_sequence
    assert not any("\n\t" in text for text in texts[:7])

    # The output should not contain a separate math-only paragraph made from
    # grouped legend symbols A/B/C.
    grouped_symbol_paras = [
        extract_math_texts(p)
        for p in out.paragraphs
        if has_math(p) and not (p.text or "").strip() and extract_math_texts(p) == ["A", "B", "C"]
    ]
    assert not grouped_symbol_paras

    print("word math symbol legend: ok")


if __name__ == "__main__":
    main()
