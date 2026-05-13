"""End-to-end test for the per-paragraph in-place translation path.

Builds a docx with TWO kinds of equation-bearing layouts and verifies each
paragraph ends up with the right English in the right place:

  A. One mixed paragraph: 'text + inline <m:oMath> + text'.
     Expected: the math element stays at its source XML position; surrounding
     text gets translated.

  B. Five-paragraph block: intro text, three standalone <m:oMath> paragraphs,
     and a legend paragraph. Expected: each standalone math paragraph stays
     untouched; the intro and legend paragraphs are translated in their own
     <w:p> elements.

Stub LLM records each call so we can see one call per text-bearing paragraph
and zero calls for pure-math paragraphs.

Run:
    uv run python scripts/test_per_paragraph.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

_ROOT = Path(__file__).resolve().parent.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from docx import Document  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402

from translate.agent.graph import build_graph  # noqa: E402
from translate.agent.state import TranslationState  # noqa: E402


def _math_run(text: str):
    r = OxmlElement("m:r")
    t = OxmlElement("m:t")
    t.text = text
    r.append(t)
    return r


def _eq(*tokens: str):
    omath = OxmlElement("m:oMath")
    for tok in tokens:
        omath.append(_math_run(tok))
    return omath


def _txt(text: str):
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    r.append(t)
    return r


def make_input(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("[발명의 설명]")
    doc.add_paragraph("[발명의 명칭]")
    doc.add_paragraph("샘플 명세서")
    doc.add_paragraph("[발명의 실시를 위한 구체적인 내용]")

    # A. Mixed paragraph: text + inline OMML + text.
    p = doc.add_paragraph()
    p._p.append(_txt("다음 식 "))
    p._p.append(_eq("A", " = ", "B", " + ", "C"))
    p._p.append(_txt("을 만족하고, 여기서 A는 두께, B는 폭, C는 높이이다."))

    # B. Five-paragraph block: intro, 3 standalone equations, legend.
    doc.add_paragraph("그리고 본 발명은 다음 식들을 만족한다:")
    p = doc.add_paragraph()
    p._p.append(_eq("X", " = ", "Y", " + ", "Z"))
    p = doc.add_paragraph()
    p._p.append(_eq("P", " = ", "Q", " * ", "R"))
    p = doc.add_paragraph()
    p._p.append(_eq("M", " = ", "N", " - ", "K"))
    doc.add_paragraph("여기서 X는 길이, Y는 면적, Z는 부피, P는 압력, Q는 온도, R은 시간이다.")

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


_CANNED = {
    "샘플 명세서": "Sample Specification",
    "다음 식": "the following equation",
    "을 만족하고": "is satisfied,",
    "A는 두께": "A is a thickness",
    "B는 폭": "B is a width",
    "C는 높이이다": "C is a height",
    "그리고 본 발명은 다음 식들을 만족한다": "And the present invention satisfies the following equations:",
    "X는 길이": "X is a length",
    "Y는 면적": "Y is an area",
    "Z는 부피": "Z is a volume",
    "P는 압력": "P is a pressure",
    "Q는 온도": "Q is a temperature",
    "R은 시간이다": "R is time",
}


class RecordingStub:
    def __init__(self):
        self.calls: list[str] = []

    def complete(self, messages):
        user = messages[1]["content"]
        for tag in ("Korean clause:\n", "Korean:\n"):
            if tag in user:
                ko = user.rsplit(tag, 1)[-1].strip()
                break
        else:
            ko = user[:80]
        self.calls.append(ko)
        for k, v in _CANNED.items():
            if k in ko:
                return json.dumps({"text": v, "key_terms": []})
        return json.dumps({"text": "(no canned)", "key_terms": []})


def _ancestors(el):
    a = el.getparent()
    while a is not None:
        yield a
        a = a.getparent()


def main() -> None:
    src = Path("data/sample_per_paragraph.docx")
    dst = Path("output/sample_per_paragraph_en.docx")
    make_input(src)
    if dst.exists():
        dst.unlink()

    stub = RecordingStub()
    initial: TranslationState = {
        "input_path": src,
        "output_path": dst,
        "font": "Times New Roman",
        "review": False,
        "verbose": False,
        "delay": 0.0,
        "client": stub,
    }
    build_graph().invoke(initial)

    print(f"Stub LLM calls: {len(stub.calls)} total")
    for i, c in enumerate(stub.calls, 1):
        print(f"  {i:>2}. {c!r}")

    print()
    print(f"=== {dst.name} ===")
    doc = Document(dst)
    for i, p in enumerate(doc.paragraphs):
        seq = []
        for el in p._p.iter():
            tag = el.tag.split("}")[-1]
            if tag == "oMath":
                # Collect math symbols to show the equation passed through.
                bits = [c.text for c in el.iter() if c.tag.endswith("}t") and c.text]
                seq.append("<EQ:" + "".join(bits) + ">")
            elif tag == "t" and el.text and not any(
                a.tag.endswith("}oMath") for a in _ancestors(el)
            ):
                seq.append(repr(el.text))
        print(f"[{i:>2}] {' | '.join(seq) if seq else '(blank)'}")


if __name__ == "__main__":
    main()
