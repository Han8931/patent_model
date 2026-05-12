"""Verify '[001]' / '[002]' paragraph IDs survive translation.

Stubs the LLM to STRIP the leading '[NNN]' from every paragraph (the exact
misbehavior the user reported). The deterministic preservation in
chunk_body + translate_body should re-attach the prefix so the output
docx keeps every '[001]' / '[002]' marker in place.

Run:
    uv run python scripts/test_paragraph_ids.py
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

_ROOT = Path(__file__).resolve().parent.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from docx import Document  # noqa: E402

from translate.agent.graph import build_graph  # noqa: E402
from translate.agent.state import TranslationState  # noqa: E402


def make_input(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("[발명의 설명]")
    doc.add_paragraph("[발명의 명칭]")
    doc.add_paragraph("반도체 패키지")
    doc.add_paragraph("[발명의 실시를 위한 구체적인 내용]")
    doc.add_paragraph("[001] 본 발명은 반도체 장치에 관한 것이다.")
    doc.add_paragraph("[002] 본 발명의 일 실시예에 따른 반도체 장치는 다음을 포함할 수 있다.")
    doc.add_paragraph("[003] 또한 본 발명의 다른 실시예는 추가적인 구성요소를 포함한다.")
    doc.add_paragraph("[0123] 4-digit IDs도 함께 처리될 수 있다.")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


_LEADING_BRACKET_RE = re.compile(r'^\[\d+\]\s*')


_CANNED = {
    "본 발명은 반도체 장치에 관한 것이다.":
        "The present invention relates to a semiconductor device.",
    "본 발명의 일 실시예에 따른 반도체 장치는 다음을 포함할 수 있다.":
        "A semiconductor device according to an embodiment of the present invention may include the following.",
    "또한 본 발명의 다른 실시예는 추가적인 구성요소를 포함한다.":
        "Another embodiment of the present invention includes additional components.",
    "4-digit IDs도 함께 처리될 수 있다.":
        "4-digit IDs may also be handled together.",
    "반도체 패키지": "Semiconductor package",
}


class IdStrippingStub:
    """Stub that simulates an LLM dropping the leading '[NNN]' prefix and
    translating only the body text. The user's reported failure mode.
    """
    def complete(self, messages):
        user = messages[1]["content"]
        korean = user.rsplit("Korean:\n", 1)[-1].strip()
        lines = [_LEADING_BRACKET_RE.sub("", line).strip() for line in korean.split("\n") if line.strip()]
        out_lines = [_CANNED.get(line, "(untranslated)") for line in lines]
        translated = " ".join(out_lines)
        return json.dumps({"text": translated, "key_terms": []})


def main() -> None:
    src = Path("data/sample_paragraph_ids.docx")
    dst = Path("output/sample_paragraph_ids_en.docx")
    src.parent.mkdir(parents=True, exist_ok=True)
    dst.parent.mkdir(parents=True, exist_ok=True)
    make_input(src)

    if dst.exists():
        dst.unlink()

    initial: TranslationState = {
        "input_path": src,
        "output_path": dst,
        "font": "Times New Roman",
        "review": False,
        "verbose": False,
        "delay": 0.0,
        "client": IdStrippingStub(),
    }
    build_graph().invoke(initial)

    print(f"\n=== {dst.name} ===")
    doc = Document(dst)
    for i, p in enumerate(doc.paragraphs):
        text = (p.text or "").strip()
        marker = "  ID-PRESERVED" if re.match(r'^\[\d{1,5}\]', text) else "             "
        print(f"[{i:>2}]{marker}  {text[:100]!r}")


if __name__ == "__main__":
    main()
