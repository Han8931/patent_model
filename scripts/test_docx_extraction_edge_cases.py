"""Verify DOCX text extraction across common Word XML edge cases.

Run:
    uv run python scripts/test_docx_extraction_edge_cases.py
"""

from __future__ import annotations

import sys
from pathlib import Path

_ROOT = Path(__file__).resolve().parent.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from docx import Document  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from translate.agent.docx_utils import extract_all_text, iter_all_paragraphs  # noqa: E402


def _run(text: str) -> OxmlElement:
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    r.append(t)
    return r


def _fld_char(kind: str) -> OxmlElement:
    r = OxmlElement("w:r")
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), kind)
    r.append(fld)
    return r


def _instr(text: str) -> OxmlElement:
    r = OxmlElement("w:r")
    instr = OxmlElement("w:instrText")
    instr.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    instr.text = text
    r.append(instr)
    return r


def _field(paragraph, instr: str, display: str) -> None:
    paragraph._p.append(_fld_char("begin"))
    paragraph._p.append(_instr(instr))
    paragraph._p.append(_fld_char("separate"))
    paragraph._p.append(_run(display))
    paragraph._p.append(_fld_char("end"))


def _tab() -> OxmlElement:
    r = OxmlElement("w:r")
    r.append(OxmlElement("w:tab"))
    return r


def _br() -> OxmlElement:
    r = OxmlElement("w:r")
    r.append(OxmlElement("w:br"))
    return r


def _no_break_hyphen() -> OxmlElement:
    r = OxmlElement("w:r")
    r.append(OxmlElement("w:noBreakHyphen"))
    return r


def main() -> None:
    doc = Document()

    p = doc.add_paragraph()
    p._p.append(_run("광 엔진 유닛(OEU)"))
    p._p.append(_tab())
    p._p.append(_run("전극(100a)"))
    p._p.append(_br())
    p._p.append(_run("비"))
    p._p.append(_no_break_hyphen())
    p._p.append(_run("일시적 매체"))
    assert extract_all_text(p) == "광 엔진 유닛(OEU)\t전극(100a)\n비-일시적 매체"

    seq = doc.add_paragraph("문단 ")
    _field(seq, " SEQ para \\* ARABIC ", "[0001]")
    seq.add_run(" 본문")
    assert extract_all_text(seq) == "문단  본문"

    ref = doc.add_paragraph("도면 ")
    _field(ref, " REF _Ref12345 \\h ", "도 1")
    ref.add_run(" 참조")
    assert extract_all_text(ref) == "도면 도 1 참조"

    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).paragraphs[0].text = "표 내부 텍스트"
    all_text = [extract_all_text(p) for p in iter_all_paragraphs(doc)]
    assert "표 내부 텍스트" in all_text

    print("docx extraction edge cases: ok")


if __name__ == "__main__":
    main()
