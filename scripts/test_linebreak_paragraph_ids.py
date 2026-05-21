"""Verify Word line breaks containing paragraph IDs survive translation.

Builds a docx where ONE Word paragraph contains BOTH '[0074]' and '[0075]'
separated by a real <w:br/>. The important part is the Word line break, not
the paragraph-ID token itself: paragraph IDs are preserved as text, and the
source <w:br/> remains a line break in the translated paragraph.

Expected output shape:

    [0074] body 1
    [0075] body 2

Run:
    uv run python scripts/test_linebreak_paragraph_ids.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

_ROOT = Path(__file__).resolve().parent.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from docx import Document  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402

from translate.agent.graph import build_graph  # noqa: E402
from translate.agent.state import TranslationState  # noqa: E402


def _t(text: str):
    """Build a w:r/<w:t> run with xml:space=preserve."""
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    r.append(t)
    return r


def _br():
    r = OxmlElement("w:r")
    r.append(OxmlElement("w:br"))
    return r


def make_input(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("[발명의 설명]")
    doc.add_paragraph("[발명의 명칭]")
    doc.add_paragraph("실시예 명세서")
    doc.add_paragraph("[발명의 실시를 위한 구체적인 내용]")

    # ONE Word paragraph with two paragraph IDs separated by a <w:br>.
    p = doc.add_paragraph()
    p._p.append(_t("[0074] 다음과 같은 표현들이 사용될 수 있다."))
    p._p.append(_br())
    p._p.append(_t("[0075] 예를 들어 이는 일 실시예에 따른 것이다."))

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


_CANNED = {
    "다음과 같은 표현들이 사용될 수 있다.": "Expressions such as the following may be used.",
    "예를 들어 이는 일 실시예에 따른 것이다.": "For example, this is according to one embodiment.",
    "실시예 명세서": "Embodiment Specification",
}


class CannedStub:
    def __init__(self):
        self.calls: list[str] = []

    def complete(self, messages):
        user = messages[1]["content"]
        for tag in ("Korean clause:\n", "Korean:\n"):
            if tag in user:
                ko = user.rsplit(tag, 1)[-1].strip()
                break
        else:
            ko = user[:80]
        self.calls.append(ko)
        for k, v in _CANNED.items():
            if k in ko:
                return json.dumps({"text": v, "key_terms": []})
        return json.dumps({"text": "(no canned)", "key_terms": []})


def main() -> None:
    src = Path("data/sample_linebreak_ids.docx")
    dst = Path("output/sample_linebreak_ids_en.docx")
    src.parent.mkdir(parents=True, exist_ok=True)
    dst.parent.mkdir(parents=True, exist_ok=True)
    make_input(src)
    if dst.exists():
        dst.unlink()

    stub = CannedStub()
    initial: TranslationState = {
        "input_path": src,
        "output_path": dst,
        "font": "Times New Roman",
        "review": False,
        "verbose": False,
        "delay": 0.0,
        "client": stub,
    }
    build_graph().invoke(initial)

    print("Stub LLM calls:")
    for i, c in enumerate(stub.calls, 1):
        print(f"  {i:>2}. {c!r}")

    print()
    print(f"=== {dst.name} ===")
    doc = Document(dst)
    for i, p in enumerate(doc.paragraphs):
        text = (p.text or "").strip()
        # Count [NNN] markers — this sample has one per source line.
        import re
        ids = re.findall(r"\[\d{1,5}\]", text)
        print(f"[{i:>2}] ids={ids}  {text!r}")


if __name__ == "__main__":
    main()
