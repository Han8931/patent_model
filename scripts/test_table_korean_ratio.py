"""Verify post-write Korean ratio sees table-wrapped Word content.

Run:
    uv run python scripts/test_table_korean_ratio.py
"""

from __future__ import annotations

import sys
from pathlib import Path

_ROOT = Path(__file__).resolve().parent.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from docx import Document  # noqa: E402

from translate.agent.nodes.write import _korean_char_ratio  # noqa: E402


def main() -> None:
    path = Path("data/sample_table_korean_ratio.docx")
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).paragraphs[0].text = "본 발명은 반도체 장치에 관한 것이다."
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)

    loaded = Document(path)
    assert len(loaded.paragraphs) == 0
    ratio = _korean_char_ratio(loaded)
    assert ratio == 1.0, ratio
    print("table korean ratio: ok")


if __name__ == "__main__":
    main()
