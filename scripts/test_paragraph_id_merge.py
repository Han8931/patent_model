"""Reproduce the user's '[0131][0132] …' merge bug.

Builds a doc where the [0131] paragraph ends with ',' (would trigger continuation
merge) and the [0132] paragraph has a stray LEADING SPACE inside its first w:t
run before the bracket. With the bug, _has_paragraph_id misses the marker so
chunk_body merges the two paragraphs and [0132]'s body text is lost in the
output.

With the fix (whitespace-tolerant regex), the two paragraphs must stay separate.

Run:
    uv run python scripts/test_paragraph_id_merge.py
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

_ROOT = Path(__file__).resolve().parent.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from docx import Document  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402

from translate.agent.graph import build_graph  # noqa: E402
from translate.agent.state import TranslationState  # noqa: E402


def _add_para_with_runs(doc, runs: list[str]):
    p = doc.add_paragraph()
    for txt in runs:
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        t.text = txt
        r.append(t)
        p._p.append(r)
    return p


def make_input(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("[발명의 설명]")
    doc.add_paragraph("[발명의 명칭]")
    doc.add_paragraph("샘플 명세서")
    doc.add_paragraph("[발명의 실시를 위한 구체적인 내용]")
    # [0131] ends with a comma — would invite continuation merge.
    _add_para_with_runs(doc, ["[0131] 첫 번째 문단의 내용,"])
    # [0132] has a LEADING SPACE before the bracket — the bug trigger.
    _add_para_with_runs(doc, [" [0132] 이미지 소유자에 대한 설명."])
    # [0133] is clean — sanity check that a normal paragraph still works.
    _add_para_with_runs(doc, ["[0133] 세 번째 문단의 내용."])
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


_CANNED = {
    "첫 번째 문단의 내용,": "The first paragraph content,",
    "이미지 소유자에 대한 설명.": "The image owner description.",
    "세 번째 문단의 내용.": "The third paragraph content.",
    "샘플 명세서": "Sample Specification",
}


class CannedStub:
    """Translate stripped body content using the canned dictionary."""
    def complete(self, messages):
        user = messages[1]["content"]
        korean = user.rsplit("Korean:\n", 1)[-1].strip()
        # Each line is a paragraph; look up in canned dict, else echo English placeholder.
        lines = [line.strip() for line in korean.split("\n") if line.strip()]
        out = [_CANNED.get(line, "(no canned translation)") for line in lines]
        return json.dumps({"text": " ".join(out), "key_terms": []})


def main() -> None:
    src = Path("data/sample_id_merge.docx")
    dst = Path("output/sample_id_merge_en.docx")
    src.parent.mkdir(parents=True, exist_ok=True)
    dst.parent.mkdir(parents=True, exist_ok=True)
    make_input(src)
    if dst.exists():
        dst.unlink()

    initial: TranslationState = {
        "input_path": src,
        "output_path": dst,
        "font": "Times New Roman",
        "review": False,
        "verbose": True,
        "delay": 0.0,
        "client": CannedStub(),
    }
    build_graph().invoke(initial)

    print(f"\n=== {dst.name} ===")
    doc = Document(dst)
    for i, p in enumerate(doc.paragraphs):
        text = (p.text or "").strip()
        ids_in_para = re.findall(r"\[\d{1,5}\]", text)
        marker = (
            "  PASS"
            if 0 <= len(ids_in_para) <= 1 else
            f"  FAIL ({len(ids_in_para)} IDs in one paragraph: {ids_in_para})"
        )
        print(f"[{i:>2}]{marker}  {text[:120]!r}")


if __name__ == "__main__":
    main()
