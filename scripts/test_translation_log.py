"""Verify PatentTranslator writes a per-file log.

Run:
    uv run python scripts/test_translation_log.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

_ROOT = Path(__file__).resolve().parent.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from docx import Document  # noqa: E402

from translate.translator import PatentTranslator  # noqa: E402


class StubClient:
    def complete(self, messages):
        user = messages[-1]["content"]
        if "반도체 장치" in user:
            return json.dumps({"text": "A semiconductor device.", "key_terms": []})
        return json.dumps({"text": "Translated text.", "key_terms": []})


def main() -> None:
    src = Path("data/sample_log.docx")
    dst = Path("output/sample_log_en.docx")
    log = Path("output/sample_log_en.log")
    doc = Document()
    doc.add_paragraph("[발명의 설명]")
    doc.add_paragraph("[발명의 명칭]")
    doc.add_paragraph("반도체 장치")
    src.parent.mkdir(parents=True, exist_ok=True)
    doc.save(src)
    for path in (dst, log):
        if path.exists():
            path.unlink()

    translator = PatentTranslator()
    translator.client = StubClient()
    translator.translate_document(
        src,
        dst,
        review=False,
        verbose=False,
        log_path=log,
        progress_callback=lambda _: None,
    )

    text = log.read_text(encoding="utf-8")
    assert "Translation log started" in text
    assert "Translation completed successfully" in text
    assert str(src) in text
    print("translation log: ok")


if __name__ == "__main__":
    main()
