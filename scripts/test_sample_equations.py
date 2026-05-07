"""Offline test for the equation-layout fix using data/sample_kr_equations.docx.

Runs the agent graph end-to-end with a stubbed LLM client — no API needed.
Two stub responses are provided:

  - "good": LLM keeps [EQUATION_1] / [EQUATION_2] interleaved with their legends.
  - "bad" : LLM groups both equations first, then both legends.

For each, we open the resulting docx and print the order of math elements vs.
text runs in the mixed paragraph, so you can see how `_replace_text_with_math_placeholders`
distributes the translation around the equations.

Run:
    uv run python scripts/test_sample_equations.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

# Make the project root importable regardless of CWD when invoked.
_ROOT = Path(__file__).resolve().parent.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from docx import Document  # noqa: E402

from translate.agent.graph import build_graph  # noqa: E402
from translate.agent.state import TranslationState  # noqa: E402


SAMPLE = Path("data/sample_kr_equations.docx")
OUT_GOOD = Path("output/sample_en_good.docx")
OUT_BAD = Path("output/sample_en_bad.docx")


GOOD_BODY = (
    "The semiconductor device satisfies [EQUATION_1], "
    "where A is a thickness, B is a width, and C is a height. "
    "The device also satisfies [EQUATION_2], "
    "where X is a length, Y is an area, and Z is a density."
)

# Anti-pattern: both markers grouped first, then both legends.
BAD_BODY = (
    "The semiconductor device satisfies [EQUATION_1] [EQUATION_2], "
    "where A is a thickness, B is a width, C is a height, "
    "X is a length, Y is an area, and Z is a density."
)


class StubClient:
    """Returns canned JSON responses regardless of input."""

    def __init__(self, body_text: str):
        self._body_text = body_text

    def complete(self, messages: list[dict]) -> str:
        # Mimic the JSON schema main code expects
        return json.dumps({"text": self._body_text, "key_terms": []})


def run(label: str, body_text: str, out_path: Path) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    if out_path.exists():
        out_path.unlink()

    initial: TranslationState = {
        "input_path": SAMPLE,
        "output_path": out_path,
        "font": "Times New Roman",
        "review": False,
        "verbose": False,
        "delay": 0.0,
        "client": StubClient(body_text),
    }
    graph = build_graph()
    graph.invoke(initial)

    print(f"\n=== {label} → {out_path} ===")
    doc = Document(out_path)
    for i, p in enumerate(doc.paragraphs):
        # Walk the XML linearly to show whether equations and text are interleaved.
        seq: list[str] = []
        for el in p._p.iter():
            tag = el.tag.split("}")[-1]
            if tag == "oMath":
                seq.append("<EQ>")
            elif tag == "t" and el.text:
                seq.append(repr(el.text))
        print(f"[{i}] {' | '.join(seq)}")


def main() -> None:
    if not SAMPLE.exists():
        raise SystemExit(
            f"missing {SAMPLE}; run `uv run python scripts/make_sample_docx.py` first"
        )
    run("GOOD (interleaved)", GOOD_BODY, OUT_GOOD)
    run("BAD  (grouped)",     BAD_BODY,  OUT_BAD)


if __name__ == "__main__":
    main()
