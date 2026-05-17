"""Regression test for ordinary '상기 ... 에서' prose.

The phrase means "in/on the above ..." in normal patent prose. It must not be
treated as an equation parameter legend such as "상기 수학식 1에서 ...".

Run:
    uv run python scripts/test_legend_false_positive.py
"""

from __future__ import annotations

import sys
from pathlib import Path

_ROOT = Path(__file__).resolve().parent.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from docx import Document  # noqa: E402

from translate.agent.nodes.chunk_body import chunk_body  # noqa: E402
from translate.agent.nodes.classify import classify  # noqa: E402
from translate.agent.paragraph_translator import (  # noqa: E402
    chunk_has_legend,
    chunk_needs_per_paragraph,
)
from translate.agent.sentence_translator import _LEGEND_HEADER_RE  # noqa: E402


BAD_PROSE = (
    "도 2의 제 5 반도체 다이 100d은 제 1 웨이퍼 구조체 WF1에 "
    "제 2 관통 비아 111의 형성 없이 형성될 수 있다. "
    "상기 제 1 웨이퍼 구조체 WF1에서 제 2 기판 101의 두께를 줄이는 "
    "그라인딩 공정도 생략될 수 있다."
)


def main() -> None:
    assert _LEGEND_HEADER_RE.search(BAD_PROSE) is None

    doc = Document()
    doc.add_paragraph("[발명의 설명]")
    doc.add_paragraph("[발명의 실시를 위한 구체적인 내용]")
    doc.add_paragraph("[0038]" + BAD_PROSE)

    state = {
        "doc": doc,
        "input_path": Path("data/published1_kr_processed.docx"),
        "progress": lambda _msg: None,
    }
    state.update(classify(state))
    state.update(chunk_body(state))

    records = state["records"]
    indexed = [None] * (max(r.index for r in records) + 1)
    for record in records:
        indexed[record.index] = record

    body_chunks = state["chunks_body"]
    assert len(body_chunks) == 1
    chunk = body_chunks[0]
    assert not chunk_has_legend(chunk, indexed)
    assert not chunk_needs_per_paragraph(chunk, indexed)

    print("legend false positive: ok")


if __name__ == "__main__":
    main()
