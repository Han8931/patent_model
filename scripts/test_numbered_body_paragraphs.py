"""Verify numbered body paragraphs such as [001]/[002] are translated.

This covers two failure modes:
  1. A [001] paragraph ending in a comma must not merge with [002].
  2. A plain-text LLM response beginning with [001] must not be rejected as JSON.

Run:
    uv run python scripts/test_numbered_body_paragraphs.py
"""

from __future__ import annotations

import sys
from pathlib import Path

_ROOT = Path(__file__).resolve().parent.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from docx import Document  # noqa: E402

from translate.agent.graph import build_graph  # noqa: E402
from translate.agent.state import TranslationState  # noqa: E402


def make_input(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("[발명의 설명]")
    doc.add_paragraph("[발명의 실시를 위한 구체적인 내용]")
    doc.add_paragraph("[001] 본 발명은 반도체 장치에 관한 것이며,")
    doc.add_paragraph("[002] 상기 반도체 장치는 기판을 포함한다.")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


class PlainTextNumberedStub:
    def complete(self, messages):
        user = messages[-1]["content"]
        korean = user.rsplit("Korean", 1)[-1]
        if "본 발명은 반도체 장치에 관한 것이며" in korean:
            return "[001] The present invention relates to a semiconductor device,"
        if "상기 반도체 장치는 기판을 포함한다" in korean:
            return "[002] the semiconductor device includes a substrate."
        return '{"text": "Translated text.", "key_terms": []}'


def main() -> None:
    src = Path("data/sample_numbered_body.docx")
    dst = Path("output/sample_numbered_body_en.docx")
    make_input(src)
    if dst.exists():
        dst.unlink()

    initial: TranslationState = {
        "input_path": src,
        "output_path": dst,
        "font": "Times New Roman",
        "review": False,
        "verbose": False,
        "delay": 0.0,
        "client": PlainTextNumberedStub(),
    }
    build_graph().invoke(initial)

    doc = Document(dst)
    texts = [(p.text or "").strip() for p in doc.paragraphs]
    numbered = [t for t in texts if t.startswith("[00")]
    assert numbered == [
        "[001] The present invention relates to a semiconductor device,",
        "[002] the semiconductor device includes a substrate.",
    ], numbered
    assert not any(any("가" <= ch <= "힣" for ch in t) for t in numbered)

    print("numbered body paragraphs: ok")
    for text in numbered:
        print(f"  {text}")


if __name__ == "__main__":
    main()
