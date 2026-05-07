"""Generate a tiny Korean sample .docx with inline OMML equations.

The middle paragraph is the interesting case for the equation-layout fix:
two inline Word equations, each followed by its own '여기서, …' legend.
A correct translation should keep them alternating (eq1 + legend1, eq2 + legend2),
not regroup them as eq1 eq2 legend1 legend2.

Run:
    uv run python scripts/make_sample_docx.py
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def _m(tag: str) -> str:
    return f"{{{_M}}}{tag}"


def _math_text(text: str) -> "OxmlElement":
    """Build an <m:r><m:t>text</m:t></m:r> wrapper used inside OMML."""
    r = OxmlElement("m:r")
    t = OxmlElement("m:t")
    t.text = text
    r.append(t)
    return r


def make_inline_equation(*tokens: str) -> "OxmlElement":
    """Build a minimal inline <m:oMath> from a sequence of literal tokens.

    Tokens are joined as separate <m:r><m:t>…</m:t></m:r> children, which is
    enough for Word to render and for our pipeline (which treats <m:oMath>
    atomically) to detect as an equation.
    """
    omath = OxmlElement("m:oMath")
    for tok in tokens:
        omath.append(_math_text(tok))
    return omath


def append_run_text(p_xml, text: str) -> None:
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    p_xml.append(r)


def main() -> None:
    out = Path(__file__).resolve().parent.parent / "data" / "sample_kr_equations.docx"
    out.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()

    # Paragraph 1 — pure text intro.
    doc.add_paragraph(
        "본 발명은 다음 두 식을 만족하는 반도체 장치에 관한 것이다."
    )

    # Paragraph 2 — mixed: text + EQ1 + legend1 + text + EQ2 + legend2.
    # This is the case the chunk_body / prompt fix is meant to keep in order.
    p2 = doc.add_paragraph()
    p2_xml = p2._p
    append_run_text(p2_xml, "상기 반도체 장치는 ")
    p2_xml.append(make_inline_equation("A", " = ", "B", " + ", "C"))
    append_run_text(
        p2_xml,
        " 을 만족하고, 여기서 A는 두께이고, B는 폭이고, C는 높이이다. 또한, ",
    )
    p2_xml.append(make_inline_equation("X", " = ", "Y", " × ", "Z"))
    append_run_text(
        p2_xml,
        " 을 만족하고, 여기서 X는 길이이고, Y는 면적이고, Z는 밀도이다.",
    )

    # Paragraph 3 — pure text closing.
    doc.add_paragraph(
        "이러한 구성을 통해 본 발명은 우수한 열적 안정성과 전기적 특성을 제공한다."
    )

    doc.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
