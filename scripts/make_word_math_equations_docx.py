"""Generate a Korean sample DOCX where math symbols use Word OMML.

This sample is intentionally stricter than the older equation samples:
  - centered equation paragraphs are Word math objects;
  - parameter symbols in the Korean legend are also inline Word math objects;
  - a claims section repeats the same structure.

Run:
    uv run python scripts/make_word_math_equations_docx.py
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement


def _math_run(text: str) -> OxmlElement:
    r = OxmlElement("m:r")
    t = OxmlElement("m:t")
    t.text = text
    r.append(t)
    return r


def _math(*tokens: str) -> OxmlElement:
    omath = OxmlElement("m:oMath")
    for token in tokens:
        omath.append(_math_run(token))
    return omath


def _add_inline_math(paragraph, *tokens: str) -> None:
    paragraph._p.append(_math(*tokens))


def _add_centered_equation(doc: Document, *tokens: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p._p.append(_math(*tokens))


def _add_legend(doc: Document) -> None:
    p = doc.add_paragraph()
    p.add_run("여기서, ")
    _add_inline_math(p, "α", "²", " + ", "1/2")
    p.add_run("는 제1 복합 가중치이고, ")
    _add_inline_math(p, "β", " + ", "γ", "/2")
    p.add_run("는 제2 복합 가중치이고, ")
    _add_inline_math(p, "γ", "²", " − ", "α")
    p.add_run("는 보정 함수이고, ")
    _add_inline_math(p, "X", "₁")
    p.add_run("는 제1 입력값이고, ")
    _add_inline_math(p, "Y", "₂")
    p.add_run("는 제2 입력값이고, ")
    _add_inline_math(p, "Z", "₃")
    p.add_run("는 제3 입력값이다.")


def main() -> None:
    out = (
        Path(__file__).resolve().parent.parent
        / "data"
        / "sample_word_math_equations.docx"
    )
    out.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    doc.add_paragraph("[발명의 설명]")
    doc.add_paragraph("[발명의 실시를 위한 구체적인 내용]")
    doc.add_paragraph("제어부는 다음 수학식들을 이용하여 출력값을 산출한다:")
    _add_centered_equation(doc, "S", "₁", " = ", "α", "X", " + ", "β")
    _add_centered_equation(doc, "S", "₂", " = ", "β", "Y", " + ", "γ")
    _add_centered_equation(doc, "S", "₃", " = ", "γ", "Z", " + ", "α")
    _add_legend(doc)

    doc.add_paragraph("[청구범위]")
    doc.add_paragraph("[청구항 1]")
    doc.add_paragraph("제어 장치로서, 다음 수학식들을 만족하는 산출부:")
    _add_centered_equation(doc, "R", "₁", " = ", "A", " + ", "B")
    _add_centered_equation(doc, "R", "₂", " = ", "B", " - ", "C")
    p = doc.add_paragraph()
    p.add_run("여기서, ")
    _add_inline_math(p, "A", "₁", " + ", "B", "₂/2")
    p.add_run("는 제1 조합 입력값이고, ")
    _add_inline_math(p, "B", "₂", " − ", "C")
    p.add_run("는 제2 조합 입력값이고, ")
    _add_inline_math(p, "C", "²", " + ", "1/2")
    p.add_run("는 보정값이다.")

    doc.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
