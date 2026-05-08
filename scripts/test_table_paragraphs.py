"""Reproduce 'translation finishes too fast' on table-wrapped patent docs.

Builds two docx files with identical text:
  1. body-level paragraphs (typical doc),
  2. same content nested inside a single root-level table cell.

Compares what python-docx's ``doc.paragraphs`` sees vs. our new
``iter_all_paragraphs`` helper. Then runs ``classify`` on the table-wrapped
doc to verify that paragraphs are now actually classified (was the symptom
the user reported).

Run:
    uv run python scripts/test_table_paragraphs.py
"""

from __future__ import annotations

import sys
from pathlib import Path

_ROOT = Path(__file__).resolve().parent.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from docx import Document  # noqa: E402

from translate.agent.docx_utils import iter_all_paragraphs  # noqa: E402
from translate.agent.nodes.classify import classify  # noqa: E402


SAMPLE_TEXT = [
    "[발명의 명칭]",
    "반도체 장치 및 그 제조 방법",
    "[기술분야]",
    "본 발명은 반도체 장치의 제조 분야에 관한 것이다.",
    "[배경기술]",
    "최근 반도체 장치의 집적도가 증가함에 따라 ...",
    "[발명의 내용]",
    "본 발명의 일 실시예에 따른 반도체 장치는 다음을 포함한다.",
    "[청구범위]",
    "[청구항 1]",
    "기판; 및 상기 기판 상의 절연막을 포함하는 반도체 장치.",
]


def make_flat(path: Path) -> None:
    doc = Document()
    for line in SAMPLE_TEXT:
        doc.add_paragraph(line)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def make_table_wrapped(path: Path) -> None:
    """Same content but inside a single 1×1 table — common in patent templates."""
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    # Replace the cell's empty paragraph with our content paragraphs.
    # Simplest: add new paragraphs to the cell.
    for i, line in enumerate(SAMPLE_TEXT):
        if i == 0:
            cell.paragraphs[0].text = line
        else:
            cell.add_paragraph(line)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def report(label: str, path: Path) -> None:
    doc = Document(path)
    flat_count = len(doc.paragraphs)
    full_count = sum(1 for _ in iter_all_paragraphs(doc))
    print(f"\n--- {label} ---")
    print(f"  doc.paragraphs (top-level only) : {flat_count}")
    print(f"  iter_all_paragraphs (with tables): {full_count}")

    state = {"doc": doc, "progress": lambda _: None}
    out = classify(state)
    records = out["records"]
    sections_seen = sorted({r.section for r in records if r.section})
    text_records = [r for r in records if r.kind == "text"]
    print(f"  classify → {len(records)} records, {len(text_records)} text, sections={sections_seen}")


def main() -> None:
    flat_path = Path("data/sample_flat.docx")
    table_path = Path("data/sample_table_wrapped.docx")
    make_flat(flat_path)
    make_table_wrapped(table_path)

    report("FLAT body paragraphs", flat_path)
    report("TABLE-WRAPPED body (this is the failing case)", table_path)


if __name__ == "__main__":
    main()
