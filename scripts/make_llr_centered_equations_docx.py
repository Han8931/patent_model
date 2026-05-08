"""Generate a Korean sample .docx for LLR/BIT parameter legend handling.

Shape:
    claim section/header
    Korean claim intro
    three centered equation paragraphs
    one combined Korean parameter legend containing LLR, theta_k, BIT values, mu, k

Run:
    uv run python scripts/make_llr_centered_equations_docx.py
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement


def _math_run(text: str) -> "OxmlElement":
    r = OxmlElement("m:r")
    t = OxmlElement("m:t")
    t.text = text
    r.append(t)
    return r


def _add_centered_equation(doc: Document, *tokens: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    omath = OxmlElement("m:oMath")
    for token in tokens:
        omath.append(_math_run(token))
    p._p.append(omath)


def main() -> None:
    out = (
        Path(__file__).resolve().parent.parent
        / "data"
        / "sample_llr_centered_equations.docx"
    )
    out.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    doc.add_paragraph("[청구범위]")
    doc.add_paragraph("[청구항 1]")
    doc.add_paragraph("복조 장치로서, 다음 식들을 만족하는 비트 신뢰도 산출부:")

    _add_centered_equation(doc, "LLR", "(", "θ", "ₖ", ")", " = ", "μ", " · ", "BIT", "₃ₖ")
    _add_centered_equation(
        doc,
        "LLR",
        "(",
        "θ",
        "ₖ",
        " + ",
        "π",
        "/4",
        ")",
        " = ",
        "μ",
        " · ",
        "BIT",
        "₃ₖ₊₁",
    )
    _add_centered_equation(
        doc,
        "LLR",
        "(",
        "θ",
        "ₖ",
        " + ",
        "π",
        "/2",
        ")",
        " = ",
        "μ",
        " · ",
        "BIT",
        "₃ₖ₊₂",
    )

    doc.add_paragraph(
        "여기서, LLR은 비트 신뢰도 데이터이고, θₖ는 k번째 위상 데이터의 위상차 값이고, "
        "BIT₃ₖ는 제1 비트 값이고, BIT₃ₖ₊₁는 제2 비트 값이고, "
        "BIT₃ₖ₊₂는 제3 비트 값이고, μ는 보정 계수이고, "
        "k는 위상 데이터의 인덱스이다."
    )

    doc.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
