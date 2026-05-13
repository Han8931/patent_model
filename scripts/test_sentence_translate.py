"""End-to-end test for the sentence-level path on an equation-bearing body chunk.

Builds a docx whose body paragraph reads:

    다음 식 [EQ A=B+C] 을 만족한다. 여기서 A는 두께, B는 폭, C는 높이이다.

A stub LLM records every call. We expect:
  * the sentence-level path to be taken (because the chunk has [EQUATION_N]),
  * one segment call for '다음 식 ', one segment call for '을 만족한다.',
    one clause call per parameter (A, B, C),
  * the rendered docx paragraph to have [EQUATION_1] in its original position,
    each parameter clause attached after it.

Run:
    uv run python scripts/test_sentence_translate.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

_ROOT = Path(__file__).resolve().parent.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from docx import Document  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402

from translate.agent.graph import build_graph  # noqa: E402
from translate.agent.state import TranslationState  # noqa: E402


def _math_run(text: str):
    r = OxmlElement("m:r")
    t = OxmlElement("m:t")
    t.text = text
    r.append(t)
    return r


def _add_para_with_inline_math(doc):
    p = doc.add_paragraph()

    # Pre-equation text
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = "다음 식 "
    r.append(t)
    p._p.append(r)

    # The equation
    omath = OxmlElement("m:oMath")
    for tok in ("A", " = ", "B", " + ", "C"):
        omath.append(_math_run(tok))
    p._p.append(omath)

    # Post-equation text + legend
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = "을 만족한다. 여기서 A는 두께, B는 폭, C는 높이이다."
    r.append(t)
    p._p.append(r)
    return p


def make_input(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("[발명의 설명]")
    doc.add_paragraph("[발명의 명칭]")
    doc.add_paragraph("실시예")
    doc.add_paragraph("[발명의 실시를 위한 구체적인 내용]")
    _add_para_with_inline_math(doc)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


# Canned per-fragment translations for the stub. Keys match exactly what the
# segment / clause translator sends (after its leading/trailing punctuation
# strip), so a substring `k in ko` test resolves them.
_CANNED = {
    "다음 식": "the following equation",
    "을 만족한다": "is satisfied",
    "A는 두께": "A is a thickness",
    "B는 폭": "B is a width",
    "C는 높이이다": "C is a height",
    "실시예": "Embodiment",
}


class RecordingStub:
    """Stub LLM that returns canned per-fragment translations and records calls."""
    def __init__(self):
        self.calls: list[str] = []

    def complete(self, messages):
        user = messages[1]["content"]
        # Pull whatever's after "Korean:\n" or "Korean clause:\n".
        for tag in ("Korean clause:\n", "Korean:\n"):
            if tag in user:
                ko = user.rsplit(tag, 1)[-1].strip()
                break
        else:
            ko = user[:80]
        self.calls.append(ko)
        # Find a canned key contained in the source.
        en = "(no canned)"
        for k, v in _CANNED.items():
            if k in ko:
                en = v
                break
        return json.dumps({"text": en, "key_terms": []})


def main() -> None:
    src = Path("data/sample_sentence_eq.docx")
    dst = Path("output/sample_sentence_eq_en.docx")
    src.parent.mkdir(parents=True, exist_ok=True)
    dst.parent.mkdir(parents=True, exist_ok=True)
    make_input(src)
    if dst.exists():
        dst.unlink()

    stub = RecordingStub()
    initial: TranslationState = {
        "input_path": src,
        "output_path": dst,
        "font": "Times New Roman",
        "review": False,
        "verbose": False,
        "delay": 0.0,
        "client": stub,
    }
    build_graph().invoke(initial)

    print("Stub LLM calls (in order):")
    for i, ko in enumerate(stub.calls, 1):
        print(f"  {i:>2}. {ko!r}")

    print()
    print(f"=== {dst.name} ===")
    doc = Document(dst)
    for i, p in enumerate(doc.paragraphs):
        seq = []
        for el in p._p.iter():
            tag = el.tag.split("}")[-1]
            if tag == "oMath":
                seq.append("<EQ>")
            elif tag == "t" and el.text and not any(
                a.tag.endswith("}oMath") for a in _ancestors(el)
            ):
                seq.append(repr(el.text))
        print(f"[{i:>2}] {' | '.join(seq) if seq else '(blank)'}")


def _ancestors(el):
    a = el.getparent()
    while a is not None:
        yield a
        a = a.getparent()


if __name__ == "__main__":
    main()
