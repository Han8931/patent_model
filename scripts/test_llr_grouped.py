"""Reproduce the user's LLR/〖BIT〗_3k case end-to-end.

Source structure mirrors what they pasted:
    [claim header]
    [intro text] ... 다음 식들을 만족하는 ... :
    [eq1]  LLR(〖BIT〗_3k) = -μ·cos(θ_k+π/8)         (centered)
    [eq2]  LLR(〖BIT〗_(3k+1)) = -μ·sin(θ_k+π/8)     (centered)
    [eq3]  LLR(〖BIT〗_(3k+2)) = |μ·sin(...)| - |μ·cos(...)|   (centered)
    [legend] 상기 수학식 2에서, LLR는 …, θ_k는 …, k는 …, 〖BIT〗_3k는 …,
             〖BIT〗_(3k+1)는 …, 〖BIT〗_(3k+2)는 …, μ는 …

Stub LLM emits the broken 'list-then-descs' output the user actually saw:
    'wherein\\nLLR θ_k 〖BIT〗_3k 〖BIT〗_(3k+1) 〖BIT〗_(3k+2) μ\\n
     is the bit reliability data; is the phase-difference value...; k is an integer'

Expected after the fix: the legend paragraph contains ONE per-clause block
where every parameter has its own '<sym> is <desc>;' clause, including the
ones that the LLM missed (placeholders).

Run:
    uv run python scripts/test_llr_grouped.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

_ROOT = Path(__file__).resolve().parent.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402

from translate.agent.graph import build_graph  # noqa: E402
from translate.agent.state import TranslationState  # noqa: E402


def _math_run(text: str):
    r = OxmlElement("m:r")
    t = OxmlElement("m:t")
    t.text = text
    r.append(t)
    return r


def _add_eq(doc, *tokens: str):
    p = doc.add_paragraph()
    omath = OxmlElement("m:oMath")
    for tok in tokens:
        omath.append(_math_run(tok))
    p._p.append(omath)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def make_input(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("[청구범위]")
    doc.add_paragraph("[청구항 1]")
    doc.add_paragraph(
        "비트 신뢰도 산출 방법으로서, 다음 수학식을 만족하는 방법:"
    )
    # 3 equations all sharing the same variables (LLR, BIT, θ, k, μ, π).
    _add_eq(doc, "LLR(", "〖BIT〗", "_3k", ") = -μ·cos(", "θ", "_k", "+π/8)")
    _add_eq(doc, "LLR(", "〖BIT〗", "_(3k+1)", ") = -μ·sin(", "θ", "_k", "+π/8)")
    _add_eq(doc, "LLR(", "〖BIT〗", "_(3k+2)", ") = |μ·sin(", "θ", "_k", "+π/8)| - |μ·cos(", "θ", "_k", "+π/8)|")
    doc.add_paragraph(
        "상기 수학식에서, LLR는 상기 비트 신뢰도 데이터이고, "
        "θ_k는 k번째 위상 차 데이터의 위상 차 값이고, "
        "k는 0 이상의 정수이고, 〖BIT〗_3k는 상기 k번째 위상의 첫 번째 비트이고, "
        "〖BIT〗_(3k+1)는 상기 k번째 위상의 두 번째 비트이고, "
        "〖BIT〗_(3k+2)는 상기 k번째 위상의 세 번째 비트이고, "
        "μ는 잡음 계수이다."
    )
    doc.save(path)


# This is the *exact* shape the user reported the LLM produced — symbols
# listed first on a separate line, then bare 'is …' fragments separated
# by ';'. Only the first three legend items got translated.
BROKEN_OUTPUT = (
    "A bit-reliability calculation method, the method comprising:\n"
    "[EQUATION_1] [EQUATION_2] [EQUATION_3] wherein\n"
    "LLR θ_k 〖BIT〗_3k 〖BIT〗_(3k+1) 〖BIT〗_(3k+2) μ\n"
    "is the bit reliability data;  is the phase-difference value of the k-th "
    "phase-difference data; k is an integer of 0 or more"
)


class StubClient:
    def complete(self, messages):
        return json.dumps({"text": BROKEN_OUTPUT, "key_terms": []})


def _alignment_of(p) -> str:
    pPr = p._p.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr")
    if pPr is None:
        return "default"
    jc = pPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}jc")
    if jc is None:
        return "default"
    return jc.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val") or "default"


def main() -> None:
    src = Path("data/sample_llr.docx")
    dst = Path("output/sample_llr_en.docx")
    src.parent.mkdir(parents=True, exist_ok=True)
    dst.parent.mkdir(parents=True, exist_ok=True)
    make_input(src)

    if dst.exists():
        dst.unlink()

    initial: TranslationState = {
        "input_path": src,
        "output_path": dst,
        "font": "Times New Roman",
        "review": False,
        "verbose": False,
        "delay": 0.0,
        "client": StubClient(),
    }
    build_graph().invoke(initial)

    doc = Document(dst)
    print(f"\n=== {dst.name} ===")
    for i, p in enumerate(doc.paragraphs):
        seq = []
        for el in p._p.iter():
            tag = el.tag.split("}")[-1]
            if tag == "oMath":
                seq.append("<EQ>")
            elif tag == "t" and el.text:
                seq.append(repr(el.text))
        print(f"[{i}] align={_alignment_of(p):<7s} {' | '.join(seq) if seq else '(blank)'}")


if __name__ == "__main__":
    main()
