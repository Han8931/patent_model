"""Print classification + first chunks for a docx file you suspect of failing.

Usage:
    uv run python scripts/diagnose_docx.py path/to/failing.docx

Output:
    - Number of paragraphs python-docx sees vs. iter_all_paragraphs sees.
    - First 30 paragraphs with their classified kind and a snippet of text.
    - Section list and per-kind counts.
    - Number of body / abstract / claims chunks the chunker produces.
    - First chunk of each section.

Paste the output into a chat to share what's happening on a failing doc.
"""

from __future__ import annotations

import sys
from pathlib import Path

_ROOT = Path(__file__).resolve().parent.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from docx import Document  # noqa: E402

from translate.agent.docx_utils import iter_all_paragraphs  # noqa: E402
from translate.agent.nodes.chunk_abstract import chunk_abstract  # noqa: E402
from translate.agent.nodes.chunk_body import chunk_body  # noqa: E402
from translate.agent.nodes.chunk_claims import chunk_claims  # noqa: E402
from translate.agent.nodes.classify import classify  # noqa: E402


def main() -> None:
    if len(sys.argv) != 2:
        print("Usage: uv run python scripts/diagnose_docx.py path/to/failing.docx")
        raise SystemExit(2)
    path = Path(sys.argv[1])
    if not path.exists():
        print(f"Not found: {path}")
        raise SystemExit(2)

    doc = Document(path)
    flat_count = len(doc.paragraphs)
    full_count = sum(1 for _ in iter_all_paragraphs(doc))
    print(f"== {path.name} ==")
    print(f"doc.paragraphs (top-level only): {flat_count}")
    print(f"iter_all_paragraphs (recursive) : {full_count}")
    if full_count == 0:
        print("(!) Document body has zero paragraphs — translation will be a no-op.")

    state = {"doc": doc, "progress": print}
    state.update(classify(state))
    records = state["records"]

    counts: dict[str, int] = {}
    for r in records:
        counts[r.kind] = counts.get(r.kind, 0) + 1
    sections = sorted({r.section for r in records if r.section})
    print(f"\nKind breakdown: {counts}")
    print(f"Sections seen : {sections or 'none-detected'}")

    print("\nFirst 30 paragraphs:")
    for r in records[:30]:
        snippet = (r.raw or "").replace("\n", " ⏎ ")[:80]
        print(f"  [{r.index:>3}] kind={r.kind:<15s} sect={str(r.section or '-'):<10s} {snippet!r}")

    state.update(chunk_body(state))
    state.update(chunk_abstract(state))
    state.update(chunk_claims(state))

    print("\nChunk counts:")
    for key in ("chunks_body", "chunks_abstract", "chunks_claims"):
        chunks = state.get(key, [])
        print(f"  {key}: {len(chunks)}")
        if chunks:
            c = chunks[0]
            preview = (c.text or "").replace("\n", " ⏎ ")[:140]
            print(f"    first chunk → {preview!r}")


if __name__ == "__main__":
    main()
