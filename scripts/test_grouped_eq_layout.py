"""End-to-end offline test for the 'N grouped equations + combined legend' fix.

Builds an in-memory docx with one claim shaped like:
    [claim header]
    [intro text] 다음 식들을 만족한다:
    [eq1]  (centered Word OMML)
    [eq2]  (centered)
    [eq3]  (centered)
    [legend text] 여기서, A는 ..., B는 ..., X는 ..., Y는 ..., M은 ..., N은 ...

For each of three stubbed LLM behaviors, runs the full agent graph and prints
the rendered docx paragraph-by-paragraph so we can verify each parameter
clause ends up adjacent to its OWN equation:

  1. interleaved   — LLM produced eq1+legend1, eq2+legend2, eq3+legend3.
  2. grouped       — LLM bunched all clauses after the last marker.
                     Variable-aware redistribution should rescue this.
  3. respectively  — LLM emitted 'A, B, ..., are X, Y, ..., respectively'.
                     postprocess() expands, then redistribution places.

Run:
    uv run python scripts/test_grouped_eq_layout.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

_ROOT = Path(__file__).resolve().parent.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402

from translate.agent.graph import build_graph  # noqa: E402
from translate.agent.state import TranslationState  # noqa: E402


def _math_run(text: str):
    r = OxmlElement("m:r")
    t = OxmlElement("m:t")
    t.text = text
    r.append(t)
    return r


def _add_equation_paragraph(doc, *tokens: str):
    p = doc.add_paragraph()
    omath = OxmlElement("m:oMath")
    for tok in tokens:
        omath.append(_math_run(tok))
    p._p.append(omath)
    return p


def make_input_docx(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("[청구범위]")
    doc.add_paragraph("[청구항 1]")
    doc.add_paragraph("반도체 장치로서, 다음 식들을 만족하는 반도체 장치:")
    for tokens in [
        ("A", " = ", "B", " + ", "C"),
        ("X", " = ", "Y", " × ", "Z"),
        ("M", " = ", "N", " - ", "P"),
    ]:
        p = _add_equation_paragraph(doc, *tokens)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        "여기서, A는 두께이고, B는 폭이고, C는 높이이고, "
        "X는 길이이고, Y는 면적이고, Z는 밀도이고, "
        "M은 질량이고, N은 부피이고, P는 압력이다."
    )
    doc.save(path)


# Three LLM behaviors — same input, very different output shapes:
INTERLEAVED = (
    "A semiconductor device satisfying:\n"
    "[EQUATION_1], where A is a thickness, B is a width, and C is a height;\n"
    "[EQUATION_2], where X is a length, Y is an area, and Z is a density; and\n"
    "[EQUATION_3], where M is a mass, N is a volume, and P is a pressure."
)
GROUPED = (
    "A semiconductor device satisfying:\n"
    "[EQUATION_1] [EQUATION_2] [EQUATION_3] where "
    "A is a thickness, B is a width, C is a height, "
    "X is a length, Y is an area, Z is a density, "
    "M is a mass, N is a volume, and P is a pressure."
)
RESPECTIVELY = (
    "A semiconductor device satisfying:\n"
    "[EQUATION_1] [EQUATION_2] [EQUATION_3] where "
    "A, B, C, X, Y, Z, M, N, P are a thickness, a width, a height, "
    "a length, an area, a density, a mass, a volume, and a pressure, respectively."
)


class StubClient:
    def __init__(self, response_text: str):
        self.response_text = response_text

    def complete(self, messages):
        return json.dumps({"text": self.response_text, "key_terms": []})


def _alignment_of(p) -> str:
    pPr = p._p.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr")
    if pPr is None:
        return "default"
    jc = pPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}jc")
    if jc is None:
        return "default"
    return jc.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val") or "default"


def inspect(path: Path, label: str) -> None:
    doc = Document(path)
    print(f"\n=== {label} → {path.name} ===")
    for i, p in enumerate(doc.paragraphs):
        seq = []
        for el in p._p.iter():
            tag = el.tag.split("}")[-1]
            if tag == "oMath":
                seq.append("<EQ>")
            elif tag == "t" and el.text:
                seq.append(repr(el.text))
        print(f"[{i}] align={_alignment_of(p):<7s} {' | '.join(seq) if seq else '(blank)'}")


def run_one(label: str, response: str, src: Path, dst: Path) -> None:
    if dst.exists():
        dst.unlink()
    initial: TranslationState = {
        "input_path": src,
        "output_path": dst,
        "font": "Times New Roman",
        "review": False,
        "verbose": False,
        "delay": 0.0,
        "client": StubClient(response),
    }
    build_graph().invoke(initial)
    inspect(dst, label)


def main() -> None:
    src = Path("data/sample_grouped_eq.docx")
    src.parent.mkdir(parents=True, exist_ok=True)
    Path("output").mkdir(parents=True, exist_ok=True)
    make_input_docx(src)

    run_one("INTERLEAVED  ",  INTERLEAVED,
            src, Path("output/sample_eq_interleaved.docx"))
    run_one("GROUPED      ",  GROUPED,
            src, Path("output/sample_eq_grouped.docx"))
    run_one("RESPECTIVELY ",  RESPECTIVELY,
            src, Path("output/sample_eq_respectively.docx"))


if __name__ == "__main__":
    main()
