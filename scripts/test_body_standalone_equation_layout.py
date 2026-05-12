"""Verify body equations stay in their original paragraph position.

Run:
    uv run python scripts/test_body_standalone_equation_layout.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

_ROOT = Path(__file__).resolve().parent.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402

from translate.agent.docx_utils import has_math  # noqa: E402
from translate.agent.graph import build_graph  # noqa: E402
from translate.agent.state import TranslationState  # noqa: E402


def _math_run(text: str):
    r = OxmlElement("m:r")
    t = OxmlElement("m:t")
    t.text = text
    r.append(t)
    return r


def _add_centered_equation(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    omath = OxmlElement("m:oMath")
    for token in ("A", " = ", "B", " + ", "C"):
        omath.append(_math_run(token))
    p._p.append(omath)


def make_input(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("[발명의 설명]")
    doc.add_paragraph("[발명의 실시를 위한 구체적인 내용]")
    doc.add_paragraph("제어부는 다음 수학식을 이용하여 값을 산출한다:")
    _add_centered_equation(doc)
    doc.add_paragraph("여기서, A는 출력값이고, B는 제1 입력값이며, C는 제2 입력값이다.")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


class EquationAwareStub:
    def complete(self, messages):
        user = messages[-1]["content"]
        if "[EQUATION_1]" in user:
            return json.dumps({
                "text": (
                    "The controller calculates a value using the following equation:\n"
                    "[EQUATION_1]\n"
                    "where A is an output value; B is a first input value; "
                    "and C is a second input value."
                ),
                "key_terms": [],
            })
        return json.dumps({"text": "Translated text.", "key_terms": []})


def main() -> None:
    src = Path("data/sample_body_standalone_equation.docx")
    dst = Path("output/sample_body_standalone_equation_en.docx")
    make_input(src)
    if dst.exists():
        dst.unlink()

    initial: TranslationState = {
        "input_path": src,
        "output_path": dst,
        "font": "Times New Roman",
        "review": False,
        "verbose": False,
        "delay": 0.0,
        "client": EquationAwareStub(),
    }
    build_graph().invoke(initial)

    doc = Document(dst)
    body = [
        ("EQ" if has_math(p) else "TEXT", (p.text or "").strip(), p.alignment)
        for p in doc.paragraphs
    ]
    interesting = [
        item for item in body
        if item[0] == "EQ" or "controller" in item[1] or item[1].startswith(("where A", "wherein A"))
    ]
    assert interesting[0][1] == "The controller calculates a value using the following equation:"
    assert interesting[1][0] == "EQ"
    assert interesting[1][2] == WD_ALIGN_PARAGRAPH.CENTER
    assert interesting[2][1].startswith(("where A is an output value;", "wherein A is an output value;"))
    assert not any(any("가" <= ch <= "힣" for ch in text) for _, text, _ in interesting)

    print("body standalone equation layout: ok")
    for kind, text, _ in interesting:
        print(f"  {kind}: {text}")


if __name__ == "__main__":
    main()
