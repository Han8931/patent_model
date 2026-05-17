"""End-to-end guard: markdown-style LLM responses are cleaned before DOCX save.

Run:
    uv run python scripts/test_no_markdown_docx_output.py
"""

from __future__ import annotations

import sys
from pathlib import Path
from types import SimpleNamespace

_ROOT = Path(__file__).resolve().parent.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from docx import Document  # noqa: E402

from translate.agent.graph import build_graph  # noqa: E402
from translate.agent.state import TranslationState  # noqa: E402


class MarkdownLeakingStub:
    def __init__(self) -> None:
        self.config = SimpleNamespace(max_tokens=4096)

    def complete(self, messages, max_tokens=None, return_meta=False):
        system = messages[0]["content"]
        user = messages[-1]["content"]
        if "patent translation reviewer" in system:
            if "Identify any quality issues" in user:
                text = '{"needs_revision": true, "issues": ["polish wording"]}'
            elif "Section: CLAIMS" in user:
                text = (
                    '[{"index": 0, "text": "**Translation**\\n\\n'
                    'Claim 1: A semiconductor package comprising: a substrate; '
                    'and a chip on the substrate."}]'
                )
            else:
                text = (
                    '[{"index": 0, "text": "**Translation**\\n\\n'
                    'A semiconductor package is provided."}]'
                )
        elif "Translate each Korean claim" in system:
            text = (
                "===== CLAIM 1 =====\n"
                "**Translation**\n\n"
                "Claim 1: A semiconductor package comprising: a substrate; "
                "and a chip on the substrate.\n\n"
                "===== GLOSSARY =====\n"
                "반도체 패키지 -> semiconductor package"
            )
        elif "abstract" in system.lower():
            text = (
                "**Abstract (USPTO style)**\n"
                "A semiconductor package is provided.\n\n"
                "===== GLOSSARY =====\n"
                "반도체 패키지 -> semiconductor package"
            )
        else:
            text = (
                "**Translation**\n\n"
                "The present disclosure relates to a **semiconductor package**.\n\n"
                "---\n\n"
                "===== GLOSSARY =====\n"
                "반도체 패키지 -> semiconductor package"
            )
        if return_meta:
            return text, "stop"
        return text


def make_input(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("[발명의 설명]")
    doc.add_paragraph("[발명의 명칭]")
    doc.add_paragraph("반도체 패키지")
    doc.add_paragraph("[발명의 실시를 위한 구체적인 내용]")
    doc.add_paragraph("[001] 본 발명은 반도체 패키지에 관한 것이다.")
    doc.add_paragraph("[특허청구범위]")
    doc.add_paragraph("【청구항 1】 반도체 패키지에 있어서, 기판 및 칩을 포함한다.")
    doc.add_paragraph("[요약서]")
    doc.add_paragraph("본 발명은 반도체 패키지를 제공한다.")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def main() -> None:
    src = Path("data/sample_no_markdown_output.docx")
    dst = Path("output/sample_no_markdown_output_en.docx")
    make_input(src)
    if dst.exists():
        dst.unlink()

    state: TranslationState = {
        "input_path": src,
        "output_path": dst,
        "font": "Times New Roman",
        "review": True,
        "verbose": False,
        "delay": 0.0,
        "client": MarkdownLeakingStub(),
    }
    build_graph().invoke(state)

    text = "\n".join(p.text or "" for p in Document(dst).paragraphs)
    assert "**" not in text
    assert "Translation" not in text
    assert "Abstract (USPTO style)" not in text
    assert "I'm ready" not in text
    assert "semiconductor package" in text
    assert "A semiconductor package comprising:" in text

    print("no markdown docx output: ok")


if __name__ == "__main__":
    main()
