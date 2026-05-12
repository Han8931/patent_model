"""Verify model 'empty text' meta-responses are not written as translations.

Run:
    uv run python scripts/test_empty_text_fallback.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

_ROOT = Path(__file__).resolve().parent.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from docx import Document  # noqa: E402

from translate.agent.glossary import clean_translation_text  # noqa: E402
from translate.agent.graph import build_graph  # noqa: E402
from translate.agent.state import TranslationState  # noqa: E402


class EmptyTextStub:
    def complete(self, messages):
        user = messages[-1]["content"]
        if "Retrying" in user or "previous response was invalid" in user:
            return json.dumps({"text": "", "key_terms": []})
        return "I got empty text, so there is nothing to translate."


def make_input(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("[발명의 설명]")
    doc.add_paragraph("[발명의 명칭]")
    doc.add_paragraph("반도체 장치")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def main() -> None:
    assert clean_translation_text("I got empty text, so there is nothing to translate.") == ""
    assert clean_translation_text("The input text is empty.") == ""
    assert clean_translation_text("[001] The present invention relates to a semiconductor device.") != ""

    src = Path("data/sample_empty_text_fallback.docx")
    dst = Path("output/sample_empty_text_fallback_en.docx")
    make_input(src)
    if dst.exists():
        dst.unlink()

    initial: TranslationState = {
        "input_path": src,
        "output_path": dst,
        "font": "Times New Roman",
        "review": False,
        "verbose": False,
        "delay": 0.0,
        "client": EmptyTextStub(),
    }
    try:
        build_graph().invoke(initial)
    except RuntimeError:
        pass

    if dst.exists():
        text = "\n".join(p.text or "" for p in Document(dst).paragraphs)
        assert "empty text" not in text.lower()
        assert "nothing to translate" not in text.lower()

    print("empty text fallback: ok")


if __name__ == "__main__":
    main()
