"""Agentic translation pipeline.

The deterministic pipeline in :mod:`translator` derives section roles from
Korean bracket markers (`[청구범위]`, `[요약서]`, etc.) and never asks the LLM
to make that decision. The agentic pipeline replaces those decisions with LLM
calls:

  1. **Section classification** — the LLM labels every paragraph as
     HEADER / CLAIM / ABSTRACT / DESCRIPTION / BLANK. Media-bearing paragraphs
     (images, equations) are detected structurally beforehand and tagged
     MEDIA so the LLM doesn't try to reclassify them from text alone.
  2. **Claim dep/ind classification** — for each detected claim, the LLM
     decides INDEPENDENT or DEPENDENT (with the first parent number for
     dependents).
  3. **Translation** — reuses :func:`translator.translate_claims` (passing
     the claim-type map so the bulk prompt sees the dep/ind layout up front),
     :func:`translator.translate_abstract`, and
     :func:`translator.translate_paragraph` with the rolling context window.
     Write-back reuses the same :mod:`docx_io` apply_* helpers.

All LLM responses are plain text — no JSON anywhere.
"""

from __future__ import annotations

import logging
import re
from collections import Counter, deque
from pathlib import Path
from typing import Callable

from docx import Document

from .client import ClientConfig, LLMClient
from .docx_io import (
    apply_abstract,
    apply_claims,
    apply_output_font,
    apply_section_headers,
    CLAIM_HEADER_RE,
    _has_media,
    collect_block,
    detect_section,
    set_paragraph_text,
)
from .prompts import (
    CLAIMS_MAX_TOKENS,
    CLAIM_CLASSIFY_PROMPT,
    DESCRIPTION_CONTEXT_WINDOW,
    SECTION_CLASSIFY_PROMPT,
)
from .translator import (
    _claim_korean_by_num,
    _setup_logger,
    _teardown_logger,
    translate_abstract,
    translate_claims,
    translate_paragraph,
)


Progress = Callable[[str], None]


# ---------------------------------------------------------------------------
# Section classification (LLM)
# ---------------------------------------------------------------------------

# Maps the uppercase LLM labels to the lowercase role names used elsewhere.
_LABEL_TO_ROLE = {
    "HEADER":      "header",
    "CLAIM":       "claim",
    "ABSTRACT":    "abstract",
    "DESCRIPTION": "description",
    "BLANK":       "blank",
    "MEDIA":       "media",
}

# Parses a classification line like "12: HEADER" or "12 - CLAIM" tolerantly.
_CLASSIFY_LINE_RE = re.compile(
    r"^\s*(?:\[)?(\d+)(?:\])?\s*[:.\-)]?\s*([A-Z]+)",
    re.IGNORECASE,
)


def _build_paragraphs_block(paragraphs: list[dict]) -> str:
    """Render the paragraph list shown to the section classifier.

    Each paragraph is shown as `[N] <first 200 chars>`. Media paragraphs are
    pre-marked so the LLM doesn't try to reason about them from text alone.
    """
    lines: list[str] = []
    for p in paragraphs:
        idx = p["index"] + 1  # 1-based for the prompt
        if p["is_media"]:
            preview = "<MEDIA paragraph: contains an image or equation>"
        else:
            text = (p["text"] or "").strip()
            if not text:
                preview = "<empty>"
            else:
                preview = text[:200] + ("…" if len(text) > 200 else "")
            # collapse newlines so multi-line paragraphs stay on one input line
            preview = preview.replace("\n", " ")
        lines.append(f"[{idx}] {preview}")
    return "\n".join(lines)


def _parse_classification(raw: str, n_paragraphs: int) -> dict[int, str]:
    """Parse the LLM's classification output into ``{0-based index: ROLE}``."""
    labels: dict[int, str] = {}
    for line in raw.splitlines():
        m = _CLASSIFY_LINE_RE.match(line)
        if not m:
            continue
        idx_1 = int(m.group(1))
        label = m.group(2).upper()
        if label not in _LABEL_TO_ROLE:
            continue
        idx_0 = idx_1 - 1
        if 0 <= idx_0 < n_paragraphs:
            labels[idx_0] = _LABEL_TO_ROLE[label]
    return labels


def classify_sections_llm(
    client: LLMClient,
    paragraphs: list[dict],
    *,
    progress: Progress = print,
    logger: logging.Logger | None = None,
) -> list[dict]:
    """Classify every paragraph via the LLM and return records.

    *paragraphs* is a list of ``{"index", "text", "is_media"}`` dicts.

    Records returned use the same shape as :func:`docx_io.scan_paragraphs`:
    ``{"index", "section", "role", "text"}`` with ``role`` ∈
    {"header", "claim", "abstract", "description", "blank", "media"}.

    Media paragraphs are forced to ``role="media"`` regardless of what the
    LLM emits. For HEADER paragraphs, the English section label is set via
    the deterministic :func:`docx_io.detect_section` so the apply step can
    rewrite the Korean bracket marker.
    """
    block = _build_paragraphs_block(paragraphs)

    raw = client.complete(
        [
            {"role": "system", "content": SECTION_CLASSIFY_PROMPT},
            {"role": "user",   "content": block},
        ],
        max_tokens=CLAIMS_MAX_TOKENS,
    )

    labels = _parse_classification(raw, n_paragraphs=len(paragraphs))

    # Defaults for any paragraphs the LLM forgot:
    #   - empty text → BLANK; otherwise → DESCRIPTION (safe pass-through).
    records: list[dict] = []
    current_section: str | None = None
    misses = 0
    for p in paragraphs:
        idx = p["index"]
        if p["is_media"]:
            role = "media"
        elif idx in labels:
            role = labels[idx]
        else:
            misses += 1
            role = "blank" if not (p["text"] or "").strip() else "description"

        rec: dict = {"index": idx, "text": p["text"], "role": role, "section": None}

        if role == "header":
            section = detect_section(p["text"])
            if section is None:
                section = "UNKNOWN SECTION"
                if logger:
                    snippet = (p["text"] or "").strip()[:80]
                    logger.warning(
                        f"LLM tagged idx={idx} as HEADER but text doesn't match any known "
                        f"section marker: '{snippet}'"
                    )
            rec["section"] = section
            current_section = section
        else:
            rec["section"] = current_section

        records.append(rec)

    if misses and logger:
        logger.warning(
            f"section classifier didn't label {misses} paragraph(s); "
            "filled defaults (blank/description)"
        )

    return records


# ---------------------------------------------------------------------------
# Claim dep/ind classification (LLM)
# ---------------------------------------------------------------------------

_CLAIM_TYPE_RE = re.compile(
    r"^\s*(\d+)\s*[:.\-)]?\s*(INDEPENDENT|DEPENDENT)\b\s*(.*)$",
    re.IGNORECASE,
)
_PARENT_RE = re.compile(r"parent\s*=\s*(\d+)", re.IGNORECASE)


def classify_claims_dep_ind_llm(
    client: LLMClient,
    korean_by_num: dict[str, str],
    *,
    logger: logging.Logger | None = None,
) -> dict[str, str]:
    """Returns ``{claim_num: "INDEPENDENT" | "DEPENDENT, parent=N"}``.

    Sends the full Korean text of every claim in one call. The text per claim
    is truncated to the first 400 chars — that's enough for the dep/ind cue
    ("제1항에 있어서…") to appear; sending more is wasted tokens.
    """
    if not korean_by_num:
        return {}

    lines = []
    for num in sorted(korean_by_num, key=lambda s: int(s)):
        text = korean_by_num[num].strip().replace("\n", " ")
        if len(text) > 400:
            text = text[:400] + "…"
        lines.append(f"[Claim {num}] {text}")
    block = "\n\n".join(lines)

    raw = client.complete(
        [
            {"role": "system", "content": CLAIM_CLASSIFY_PROMPT},
            {"role": "user",   "content": block},
        ],
        max_tokens=2048,
    )

    types: dict[str, str] = {}
    for line in raw.splitlines():
        m = _CLAIM_TYPE_RE.match(line)
        if not m:
            continue
        num = m.group(1)
        kind = m.group(2).upper()
        rest = m.group(3)
        if num not in korean_by_num:
            continue
        if kind == "INDEPENDENT":
            types[num] = "INDEPENDENT"
        else:
            pm = _PARENT_RE.search(rest)
            types[num] = f"DEPENDENT, parent={pm.group(1)}" if pm else "DEPENDENT"

    if logger:
        missing = sorted(set(korean_by_num) - set(types), key=lambda s: int(s))
        if missing:
            logger.warning(
                f"claim classifier didn't label claim(s) {missing}; "
                "they will be sent to the translator without a type hint"
            )

    return types


# ---------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------

def translate_file_agent(
    in_path: Path,
    out_path: Path,
    *,
    progress: Progress = print,
    model: str | None = None,
) -> None:
    """Agentic pipeline: LLM-classified sections + LLM dep/ind claims.

    Writes the same per-file ``<out_path>.log`` as the heuristic pipeline.
    The translation, glossary, and write-back logic is shared with
    :func:`translator.translate_file`; only the classification stages are new.
    """
    out_path.parent.mkdir(parents=True, exist_ok=True)
    log_path = out_path.with_suffix(".log")
    logger = _setup_logger(in_path, log_path)
    logger.info(f"AGENT translate {in_path} -> {out_path}")

    try:
        config = ClientConfig.from_env()
        if model:
            config.model = model
        logger.info(f"model={config.model} base_url={config.base_url}")
        client = LLMClient(config)
        doc = Document(in_path)

        # Pre-scan: just text + structural media flag. No section/role yet.
        paragraphs = [
            {"index": i, "text": p.text, "is_media": _has_media(p)}
            for i, p in enumerate(doc.paragraphs)
        ]

        # ===== agent step 1: LLM section classification =====
        progress(f"[agent] classifying {len(paragraphs)} paragraphs by section…")
        try:
            records = classify_sections_llm(client, paragraphs, progress=progress, logger=logger)
        except Exception as e:
            progress(f"      ERROR: section classification failed — {e!r}")
            logger.exception("section classification failed entirely")
            raise

        roles = Counter(r["role"] for r in records)
        progress(f"      roles: {dict(roles)}")
        logger.info(f"role counts: {dict(roles)}")

        # ===== agent step 2: LLM claim dep/ind classification =====
        korean_by_num = _claim_korean_by_num(records)
        claim_types: dict[str, str] = {}
        if korean_by_num:
            progress(f"[agent] classifying {len(korean_by_num)} claim(s) by dep/ind…")
            try:
                claim_types = classify_claims_dep_ind_llm(client, korean_by_num, logger=logger)
                logger.info(f"claim types: {claim_types}")
                ind = [n for n, t in claim_types.items() if t == "INDEPENDENT"]
                dep = [n for n, t in claim_types.items() if t.startswith("DEPENDENT")]
                progress(f"      independent={ind} dependent={dep}")
            except Exception as e:
                progress(f"      WARNING: claim classification failed — {e!r}; continuing")
                logger.exception("claim dep/ind classification failed; continuing without type info")

        # ===== 1) claims translation (with type annotation) =====
        claims_block = collect_block(records, "claim")
        claims_en: dict[str, str] = {}
        glossary: dict[str, str] = {}
        if claims_block:
            n_paras = sum(1 for r in records if r["role"] == "claim")
            progress(f"[1/3] Translating claims ({n_paras} paragraphs, single LLM call)…")
            try:
                claims_en, glossary = translate_claims(
                    client, claims_block, records,
                    progress=progress, logger=logger,
                    claim_types=claim_types or None,
                )
                progress(f"      got {len(claims_en)} claim(s); glossary={len(glossary)} entries")
                logger.info(f"claims translated: {len(claims_en)}; glossary entries: {len(glossary)}")
            except Exception as e:
                progress(f"      ERROR: claims translation failed — {e!r}")
                logger.exception("claims translation failed entirely")
        else:
            progress("[1/3] No claims section found — skipping.")
            logger.info("no claims section")

        # ===== 2) abstract =====
        abstract_block = collect_block(records, "abstract")
        abstract_en = ""
        if abstract_block:
            progress("[2/3] Translating abstract…")
            try:
                abstract_en = translate_abstract(client, abstract_block, glossary)
            except Exception as e:
                progress(f"      ERROR: abstract translation failed — {e!r}")
                logger.exception("abstract translation failed")
        else:
            progress("[2/3] No abstract section found — skipping.")
            logger.info("no abstract section")

        # ===== 3) description (rolling context) =====
        desc_records = [r for r in records if r["role"] == "description" and r["text"].strip()]
        progress(
            f"[3/3] Translating description ({len(desc_records)} paragraphs, "
            f"context window={DESCRIPTION_CONTEXT_WINDOW})…"
        )
        total = len(desc_records)
        failures = 0
        context: deque = deque(maxlen=DESCRIPTION_CONTEXT_WINDOW)
        for n, r in enumerate(desc_records, 1):
            try:
                translation = translate_paragraph(
                    client, r["text"], glossary, context=context
                )
                r["translation"] = translation
                context.append((r["text"], translation))
            except Exception as e:
                failures += 1
                snippet = r["text"][:160].replace("\n", " ")
                logger.error(
                    f"description paragraph idx={r['index']} failed: {e!r}; korean='{snippet}…'"
                )
            if n % 10 == 0 or n == total:
                progress(f"      {n}/{total}")
        if failures:
            progress(f"      WARNING: {failures} paragraph(s) failed; Korean left in place (see {log_path.name})")
            logger.warning(f"{failures} description paragraph(s) failed to translate")

        # ===== write back =====
        apply_section_headers(doc, records)
        apply_claims(doc, records, claims_en, progress=progress)
        apply_abstract(doc, records, abstract_en)
        for r in desc_records:
            translation = r.get("translation")
            if translation:
                set_paragraph_text(doc.paragraphs[r["index"]], translation)
            # else: paragraph failed — Korean stays so reviewer sees it

        apply_output_font(doc)

        doc.save(out_path)
        progress(f"wrote {out_path}")
        logger.info(f"wrote {out_path}")
    except Exception:
        logger.exception("translate_file_agent failed")
        raise
    finally:
        _teardown_logger(logger)
