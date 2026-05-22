"""docx manipulation helpers — pure functions, no state."""

from __future__ import annotations

import hashlib
import re
import unicodedata
from copy import deepcopy

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph

try:
    from lxml import etree as _etree
except ImportError:  # python-docx pulls lxml, but be defensive.
    import xml.etree.ElementTree as _etree


_W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
_M = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
_XML_SPACE = '{http://www.w3.org/XML/1998/namespace}space'
_HANGUL_RE = re.compile(r'[가-힯]')
_EQUATION_PLACEHOLDER = '[EQUATION]'
_EQUATION_TOKEN_RE = re.compile(r'\[EQUATION(?:_\d+)?\]')
_EQUATION_PLACEHOLDER_RE = re.compile(r'\s*\[EQUATION(?:_\d+)?\]\s*')
_FORMULA_OPERATOR_RE = re.compile(r'[=<>≤≥≈∑∫√]|\b(?:sin|cos|tan|log|ln|exp)\b', re.IGNORECASE)
_DEFAULT_FONT_SIZE_PT = 12


def _get_or_add(parent, child_qname: str):
    child = parent.find(qn(child_qname))
    if child is None:
        child = OxmlElement(child_qname)
        if child_qname == 'w:rPr' and parent.tag == qn('w:r'):
            parent.insert(0, child)
        else:
            parent.append(child)
    return child


def _apply_rpr_font(rpr, font_name: str, font_size_pt: int | float = _DEFAULT_FONT_SIZE_PT) -> None:
    r_fonts = _get_or_add(rpr, 'w:rFonts')
    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        r_fonts.set(qn(attr), font_name)

    half_points = str(int(round(float(font_size_pt) * 2)))
    sz = _get_or_add(rpr, 'w:sz')
    sz.set(qn('w:val'), half_points)
    sz_cs = _get_or_add(rpr, 'w:szCs')
    sz_cs.set(qn('w:val'), half_points)


def set_run_font(run, font_name: str, font_size_pt: int | float = _DEFAULT_FONT_SIZE_PT) -> None:
    """Set all Word font slots and size on a python-docx run."""
    rpr = _get_or_add(run._r, 'w:rPr')
    _apply_rpr_font(rpr, font_name, font_size_pt)


def _set_xml_run_font(r_el, font_name: str, font_size_pt: int | float) -> None:
    rpr = _get_or_add(r_el, 'w:rPr')
    _apply_rpr_font(rpr, font_name, font_size_pt)


def _set_xml_paragraph_mark_font(p_el, font_name: str, font_size_pt: int | float) -> None:
    ppr = _get_or_add(p_el, 'w:pPr')
    rpr = _get_or_add(ppr, 'w:rPr')
    _apply_rpr_font(rpr, font_name, font_size_pt)


def normalize_document_font(
    doc,
    font_name: str = "Times New Roman",
    font_size_pt: int | float = _DEFAULT_FONT_SIZE_PT,
) -> None:
    """Force a consistent font family and size across the generated DOCX.

    Word stores fonts in separate slots for ASCII, East Asian, and complex
    scripts. Setting only ``run.font.name`` leaves some documents visually
    inconsistent, especially when the source DOCX had Korean template styles.
    """
    for style in doc.styles:
        font = getattr(style, "font", None)
        if font is None:
            continue
        font.name = font_name
        font.size = Pt(font_size_pt)
        rpr = getattr(getattr(style, "element", None), "rPr", None)
        if rpr is not None:
            _apply_rpr_font(rpr, font_name, font_size_pt)

    seen: set[int] = set()
    parts = [doc.part]
    parts.extend(getattr(doc.part.package, "parts", []))
    for part in parts:
        root = getattr(part, 'element', None)
        if root is None:
            continue
        marker = id(root)
        if marker in seen:
            continue
        seen.add(marker)
        for p_el in root.iter(f'{{{_W}}}p'):
            _set_xml_paragraph_mark_font(p_el, font_name, font_size_pt)
        for r_el in root.iter(f'{{{_W}}}r'):
            _set_xml_run_font(r_el, font_name, font_size_pt)


def iter_all_paragraphs(doc):
    """Yield every <w:p> in document body in source order, INCLUDING those
    nested inside tables (and tables-in-tables).

    python-docx's ``doc.paragraphs`` only returns paragraphs that are direct
    children of <w:body>; content placed inside a <w:tbl> is invisible to it.
    Many patent templates wrap the entire description / claims / abstract in
    a single root-level table — for those documents ``doc.paragraphs`` returns
    almost nothing and translation finishes in milliseconds with the source
    text untouched. This iterator walks the body subtree so the classifier
    sees every paragraph regardless of container.

    The yielded objects are real ``docx.text.paragraph.Paragraph`` instances,
    so existing helpers (``replace_text``, ``insert_para_after``, run access)
    work on them unchanged.
    """
    body = doc.element.body
    P_TAG = f'{{{_W}}}p'
    for p_el in body.iter(P_TAG):
        yield Paragraph(p_el, body)


def has_drawing(para) -> bool:
    """True iff the paragraph contains a real <w:drawing> (image/figure)."""
    return para._p.find('.//{%s}drawing' % _W) is not None


def has_math(para) -> bool:
    """True iff the paragraph contains an OMML equation."""
    return (
        para._p.find('.//{%s}oMath' % _M) is not None
        or para._p.find('.//{%s}oMathPara' % _M) is not None
    )


def has_non_text_content(para) -> bool:
    """Backwards-compat alias — true if the paragraph has math OR drawing."""
    return has_drawing(para) or has_math(para)


def remove_paragraph(para) -> None:
    """Remove a paragraph from the document XML."""
    p = para._p
    parent = p.getparent()
    if parent is not None:
        parent.remove(p)


def text_runs(para) -> list:
    return [r for r in para.runs if r.text]


def extract_all_text(para, *, math_as_placeholder: bool = False) -> str:
    """Concatenate text from <w:t> and translatable <m:t> in document order.

    Korean equation paragraphs often embed Korean labels or 'where ...' clauses
    inside <m:t> elements; para.text only returns <w:t> content and misses them.
    Formula-only Word equations are exposed as [EQUATION] placeholders instead
    of raw math symbols so the LLM does not translate or duplicate the formula.
    Use a w:br as a soft separator to mirror Word's visual line breaks.

    ``math_as_placeholder``: when True, ALWAYS render an <m:oMath> element as
    ``[EQUATION]`` regardless of its length / inline-symbol heuristic. The
    per-paragraph translator uses this so the writer can interleave text
    around the math elements at their original XML position; without it,
    short OMML symbols (like ``E_k`` or ``x²``) get inlined as raw text,
    no marker is emitted, and ``_replace_text_with_math_placeholders``
    can't run, so the math drifts to the end of the paragraph.
    """
    parts: list[str] = []
    _append_translatable_text(para._p, parts, math_as_placeholder=math_as_placeholder)
    return ''.join(parts)


def _local_name(el) -> str:
    tag = el.tag
    return tag.split('}')[-1] if '}' in tag else tag


def _is_math_element(el) -> bool:
    return el.tag in (f'{{{_M}}}oMath', f'{{{_M}}}oMathPara')


def _element_text(el) -> str:
    return ''.join(
        child.text or ''
        for child in el.iter()
        if _local_name(child) == 't'
    )


def _is_inline_math_symbol(text: str) -> bool:
    """True for short Word-math expressions that should remain in text.

    Parameter legends often encode symbols/expressions such as A, B, C, α,
    θ_k, 〖BIT〗_(3k+1), or α² + 1/2 as tiny OMML objects. Treating those as
    [EQUATION] destroys the symbol-description pairing. Full formulas with
    equation/comparison operators still become placeholders.
    """
    s = text.strip()
    if not s:
        return False
    if _FORMULA_OPERATOR_RE.search(s):
        return False
    compact = re.sub(r'\s+', '', s)
    return len(compact) <= 24


def _append_translatable_text(
    el,
    parts: list[str],
    state: dict | None = None,
    *,
    math_as_placeholder: bool = False,
) -> None:
    """Append paragraph text while treating each top-level equation as atomic.

    ``state`` carries a field stack across recursion. Field instruction text
    is never visible text. Field display text is visible and should be
    translated for fields such as HYPERLINK/REF, but SEQ-style auto-numbering
    fields are skipped so generated paragraph numbers do not leak into chunks.
    """
    if state is None:
        state = {"fields": []}

    if _is_math_element(el):
        math_text = _element_text(el)
        if not math_text:
            return
        if _HANGUL_RE.search(math_text):
            parts.append(math_text)
        elif math_as_placeholder:
            # Per-paragraph translator path: force a placeholder so the
            # writer can interleave text around the math element.
            parts.append(_EQUATION_PLACEHOLDER)
        elif _is_inline_math_symbol(math_text):
            parts.append(math_text)
        else:
            parts.append(_EQUATION_PLACEHOLDER)
        return

    if _is_inline_image_run(el):
        # Image-equation: <w:r><w:drawing>...</w:drawing></w:r>. Treat the
        # whole run as one [EQUATION] marker so the LLM keeps a placeholder
        # in the right relative position and the writer can reinsert text
        # around the run.
        if not _skip_current_field_display(state):
            parts.append(_EQUATION_PLACEHOLDER)
        return

    local = _local_name(el)
    if local == 'simpleField':
        instr = el.get(qn('w:instr')) or ''
        if _field_instruction_skips_display(instr):
            return
        for child in el:
            _append_translatable_text(
                child, parts, state,
                math_as_placeholder=math_as_placeholder,
            )
        return
    if local == 'fldChar':
        # w:fldChar attribute is in the WordprocessingML namespace.
        ftype = el.get(qn('w:fldCharType'))
        if ftype == 'begin':
            state["fields"].append({"instr": "", "phase": "instr", "skip": False})
        elif ftype == 'separate' and state["fields"]:
            field = state["fields"][-1]
            field["phase"] = "display"
            field["skip"] = _field_instruction_skips_display(field["instr"])
        elif ftype == 'end' and state["fields"]:
            state["fields"].pop()
        return
    if local == 'instrText':
        # Field instruction text — never visible content.
        if state["fields"]:
            state["fields"][-1]["instr"] += el.text or ''
        return
    if local == 't':
        if el.text and not _skip_current_field_display(state):
            parts.append(el.text)
        return
    if local in ('br', 'cr'):
        parts.append('\n')
        return
    if local == 'tab':
        parts.append('\t')
        return
    if local == 'noBreakHyphen':
        parts.append('-')
        return
    if local == 'softHyphen':
        return
    if local in ('delText',):
        return

    for child in el:
        _append_translatable_text(
            child, parts, state,
            math_as_placeholder=math_as_placeholder,
        )


def _field_instruction_skips_display(instr: str) -> bool:
    """True for generated field displays that should not become source text."""
    head = (instr or "").strip().split(maxsplit=1)
    if not head:
        return False
    return head[0].upper() in {"SEQ", "PAGE", "NUMPAGES"}


def _skip_current_field_display(state: dict) -> bool:
    for field in state.get("fields", []):
        if field.get("phase") == "instr":
            return True
        if field.get("skip"):
            return True
    return False


def _top_level_math_elements(para) -> list:
    elements: list = []

    def visit(el, inside_math: bool = False) -> None:
        is_math = _is_math_element(el)
        if is_math and not inside_math:
            elements.append(el)
            return
        for child in el:
            visit(child, inside_math or is_math)

    visit(para._p)
    return elements


def _is_inline_image_run(el) -> bool:
    """True iff ``el`` is a <w:r> run containing an inline <w:drawing>.

    Image-equations in patents are commonly stored as PNG/JPG inside a
    <w:r><w:drawing>...</w:drawing></w:r>. Treating these runs as opaque
    atoms (same as <m:oMath>) lets us emit [EQUATION] markers for them and
    preserve their position when text is written back.
    """
    if _local_name(el) != 'r':
        return False
    return el.find('.//{%s}drawing' % _W) is not None


def _top_level_inline_atoms(para) -> list:
    """Top-level positional atoms in document order: <m:oMath> and
    <m:oMathPara> for OMML equations, plus <w:r> wrappers holding inline
    <w:drawing> images for image-based equations.

    Each returned element becomes an anchor for text insertion: surrounding
    text is added via ``addprevious`` / ``addnext`` so the atom never moves.
    """
    atoms: list = []

    def visit(el, inside_atom: bool = False) -> None:
        if _is_math_element(el):
            if not inside_atom:
                atoms.append(el)
            return  # never recurse into math
        if _is_inline_image_run(el):
            if not inside_atom:
                atoms.append(el)
            return  # never recurse into the run holding a drawing
        for child in el:
            visit(child, inside_atom)

    visit(para._p)
    return atoms


def _omath_xml_signature(el) -> str:
    """Short hash of an OMML element's full XML — stable identity across moves."""
    try:
        xml = _etree.tostring(el, method="xml")
    except Exception:
        # Fallback: text content only.
        xml = _element_text(el).encode("utf-8", "ignore")
    return hashlib.md5(xml).hexdigest()[:10]


def snapshot_math_locations(doc) -> list[dict]:
    """Capture (signature, paragraph_index, short text) for every <m:oMath>
    element in the document, in document order.

    Used as a baseline taken right after ``load`` so we can detect at write
    completion whether any equation got LOST, MOVED to a different paragraph,
    or DUPLICATED. Signatures are 10-char md5 prefixes of the OMML XML, so
    the integrity report can name an equation by signature + short text
    snippet without exposing the full source content of the document.
    """
    out: list[dict] = []
    for idx, p in enumerate(iter_all_paragraphs(doc)):
        for math in _top_level_math_elements(p):
            out.append({
                "sig": _omath_xml_signature(math),
                "src_idx": idx,
                "text": _element_text(math)[:40],
            })
    return out


def verify_math_integrity(doc, snapshot: list[dict]) -> dict:
    """Compare the current document's math elements against ``snapshot``.

    Returns a dict::

        {
            "ok":         bool,
            "lost":       [snapshot rows missing from the output],
            "duplicated": [list of signatures present more than once now],
            "moved":      [(snapshot row, new paragraph index)],
        }
    """
    if not snapshot:
        return {"ok": True, "lost": [], "duplicated": [], "moved": []}

    current: dict[str, list[int]] = {}
    for idx, p in enumerate(iter_all_paragraphs(doc)):
        for math in _top_level_math_elements(p):
            current.setdefault(_omath_xml_signature(math), []).append(idx)

    lost: list[dict] = []
    duplicated: list[str] = []
    moved: list[tuple[dict, int]] = []
    for orig in snapshot:
        sig = orig["sig"]
        if sig not in current:
            lost.append(orig)
            continue
        cur_indices = current[sig]
        if len(cur_indices) > 1:
            duplicated.append(sig)
        cur_idx = cur_indices[0]
        if cur_idx != orig["src_idx"]:
            moved.append((orig, cur_idx))

    ok = not lost and not duplicated and not moved
    return {"ok": ok, "lost": lost, "duplicated": list(set(duplicated)), "moved": moved}


def extract_math_texts(para) -> list[str]:
    """Return visible text from each top-level equation in a paragraph.

    Unlike extract_all_text(), this exposes formula notation so prompts can use
    it as non-output context for assigning parameter legends to equations.
    """
    return [
        text for text in (
            _element_text(el).strip()
            for el in _top_level_math_elements(para)
        )
        if text
    ]


def _remove_korean_math(para) -> None:
    """Remove equation XML whose own text contains Korean.

    Formula-only equations are preserved. Equations containing Korean are removed
    after the English translation is written; otherwise Word-equation text stored
    in <m:t> remains visible because python-docx text runs do not own it.
    """
    for el in list(_top_level_math_elements(para)):
        if not _HANGUL_RE.search(_element_text(el)):
            continue
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)


def _strip_equation_placeholders(text: str) -> str:
    return _EQUATION_PLACEHOLDER_RE.sub(' ', text).strip()


def consolidate_formula_math_into(target_para, source_paras) -> int:
    """Detach formula <m:oMath> from each source paragraph and append to target.

    Used when a translation chunk spans multiple paragraphs but the equations
    live in trailing paragraphs. Moving them into the head paragraph lets
    _replace_text_with_math_placeholders interleave the translation around
    them inside a single paragraph, keeping equations visually aligned with
    the English text that references them.

    Equations whose text contains Hangul are NOT moved — they will be removed
    by _remove_korean_math during the head's replace_text call (their Korean
    content is replaced by the translation in the head).

    Returns the number of equations moved.
    """
    moved = 0
    target_p = target_para._p
    for src_para in source_paras:
        for math_el in list(_top_level_math_elements(src_para)):
            if _HANGUL_RE.search(_element_text(math_el)):
                continue  # Korean math will be removed elsewhere
            parent = math_el.getparent()
            if parent is not None:
                parent.remove(math_el)
            target_p.append(math_el)
            moved += 1
    return moved


def _append_line_to_run(r, part: str) -> None:
    """Append one logical line of text to an existing <w:r>.

    Leading tab characters become real ``<w:tab/>`` elements so Word
    respects tab stops; the remainder goes into ``<w:t>`` with
    ``xml:space="preserve"`` whenever it has leading or trailing spaces
    so the XML parser doesn't collapse them.
    """
    tab_count = 0
    while part.startswith('\t'):
        tab_count += 1
        part = part[1:]
    for _ in range(tab_count):
        r.append(OxmlElement('w:tab'))
    t = OxmlElement('w:t')
    if part.startswith(' ') or part.endswith(' '):
        t.set(_XML_SPACE, 'preserve')
    t.text = part
    r.append(t)


def write_run_with_breaks(
    run,
    text: str,
    font_name: str,
    font_size_pt: int | float = _DEFAULT_FONT_SIZE_PT,
) -> None:
    """Replace run content; convert \\n into <w:br/> and leading \\t into <w:tab/>."""
    r = run._r
    for child in list(r):
        local = child.tag.split('}')[-1] if '}' in child.tag else child.tag
        if local in ('t', 'br', 'tab'):
            r.remove(child)

    parts = text.split('\n')
    for idx, part in enumerate(parts):
        _append_line_to_run(r, part)
        if idx < len(parts) - 1:
            r.append(OxmlElement('w:br'))

    set_run_font(run, font_name, font_size_pt)


def _build_text_run(
    text: str,
    font_name: str,
    font_size_pt: int | float = _DEFAULT_FONT_SIZE_PT,
):
    new_r = OxmlElement('w:r')
    new_rpr = OxmlElement('w:rPr')
    _apply_rpr_font(new_rpr, font_name, font_size_pt)
    new_r.append(new_rpr)

    parts = text.split('\n')
    for idx, part in enumerate(parts):
        _append_line_to_run(new_r, part)
        if idx < len(parts) - 1:
            new_r.append(OxmlElement('w:br'))
    return new_r


def _replace_text_with_math_placeholders(
    para,
    new_text: str,
    font_name: str,
    font_size_pt: int | float = _DEFAULT_FONT_SIZE_PT,
) -> bool:
    """Replace paragraph text while keeping every inline atom at its source
    XML position.

    Atoms are <m:oMath> elements AND <w:r> runs holding <w:drawing> (image-
    based equations). The new_text is expected to carry one [EQUATION] /
    [EQUATION_N] marker per atom; the markers split the translation into
    segments that are inserted via ``addprevious`` / ``addnext`` around each
    atom. Math elements containing Hangul labels are excluded (those are
    handled by ``_remove_korean_math``); image-runs are always included.
    """
    atoms: list = []
    only_math: list = []  # subset used for variable-aware redistribution
    for el in _top_level_inline_atoms(para):
        if _is_math_element(el):
            if _HANGUL_RE.search(_element_text(el)):
                continue
            only_math.append(el)
            atoms.append(el)
        else:
            atoms.append(el)  # image-equation run — no variables, no filter

    if not atoms or not _EQUATION_TOKEN_RE.search(new_text):
        return False

    parts = _EQUATION_TOKEN_RE.split(new_text, maxsplit=len(atoms))
    if len(parts) < 2:
        return False
    while len(parts) < len(atoms) + 1:
        parts.append('')
    parts = [_EQUATION_TOKEN_RE.sub('', part) for part in parts]

    # When the LLM grouped every [EQUATION_N] marker together so all the
    # parameter clauses ended up in one slot, parse the description and
    # reassign clauses to math elements by symbol matching. Lazy import to
    # break the docx_utils ↔ nodes.write cycle. Only math elements carry
    # variables — image-runs don't, so redistribution applies only when at
    # least two math elements are present.
    if len(only_math) >= 2 and len(only_math) == len(atoms):
        from .nodes.write import _redistribute_inline_math
        parts = _redistribute_inline_math(parts, only_math)

    # Clear existing normal text. Atom elements (math XML, drawing runs)
    # remain in their original spots so positioning is preserved.
    atom_runs = {a for a in atoms if _is_inline_image_run(a)}
    for run in para.runs:
        if run._r in atom_runs:
            continue  # don't blank the run that holds the image equation
        write_run_with_breaks(run, '', font_name, font_size_pt)

    for idx, anchor in enumerate(atoms):
        before = parts[idx]
        if before:
            anchor.addprevious(_build_text_run(before, font_name, font_size_pt))

    after = parts[len(atoms)]
    if after:
        atoms[-1].addnext(_build_text_run(after, font_name, font_size_pt))
    return True


def replace_text(
    para,
    new_text: str,
    font_name: str,
    font_size_pt: int | float = _DEFAULT_FONT_SIZE_PT,
) -> None:
    """Write new_text into the paragraph's text runs.

    Formula-only equations are preserved. Equations containing Korean text are
    removed after their English translation is written; otherwise the original
    Korean <m:t> text remains visible because python-docx text runs do not own it.
    """
    _remove_korean_math(para)
    if _replace_text_with_math_placeholders(para, new_text, font_name, font_size_pt):
        return

    new_text = _strip_equation_placeholders(new_text)
    runs = text_runs(para)
    if not runs:
        if not new_text:
            return
        _prepend_text_run(para, new_text, font_name, font_size_pt)
        return
    write_run_with_breaks(runs[0], new_text, font_name, font_size_pt)
    for run in runs[1:]:
        run.text = ''
        set_run_font(run, font_name, font_size_pt)
    for run in para.runs:
        set_run_font(run, font_name, font_size_pt)


def _prepend_text_run(
    para,
    text: str,
    font_name: str,
    font_size_pt: int | float = _DEFAULT_FONT_SIZE_PT,
) -> None:
    """Insert a new <w:r><w:t>...</w:t></w:r> as the first child of <w:p>."""
    new_r = _build_text_run(text, font_name, font_size_pt)

    # Insert after <w:pPr> if present, otherwise at the start
    p = para._p
    pPr = p.find(qn('w:pPr'))
    if pPr is not None:
        pPr.addnext(new_r)
    else:
        p.insert(0, new_r)


def insert_para_after(
    ref_para,
    text: str,
    font_name: str,
    *,
    format_ref_para=None,
    font_size_pt: int | float = _DEFAULT_FONT_SIZE_PT,
):
    """Insert a new paragraph immediately after ``ref_para``.

    The ``text`` may contain ``\\n`` characters from ``postprocess``'s
    semicolon-and-indent expansion. Those become real ``<w:br/>`` elements
    and per-line ``<w:t xml:space="preserve">`` runs — without this split,
    XML treats the raw newline as whitespace and Word collapses it back
    into a single visual line, losing the indent.

    If ``format_ref_para`` is given, clone its <w:pPr> onto the new paragraph
    so the inserted paragraph inherits alignment, indentation, spacing, and
    style. This matters when the new paragraph sits between centered equation
    paragraphs and a left-aligned legend — without a format reference, the
    new paragraph uses Word defaults and visually clashes.
    """
    new_p = OxmlElement('w:p')

    if format_ref_para is not None:
        ref_pPr = format_ref_para._p.find(qn('w:pPr'))
        if ref_pPr is not None:
            new_p.append(deepcopy(ref_pPr))

    # _build_text_run already handles the \n → <w:br/> split with
    # xml:space="preserve" on lines that have leading/trailing whitespace.
    new_p.append(_build_text_run(text, font_name, font_size_pt))
    ref_para._p.addnext(new_p)
    return Paragraph(new_p, ref_para._parent)


def word_count(text: str) -> int:
    return len(text.split())


# ---------------------------------------------------------------------------
# Equation variable extraction
# ---------------------------------------------------------------------------
# Used by the claim writer to redistribute a grouped parameter legend across
# its equations. We extract identifier-like tokens from <m:oMath> XML and
# normalize math-italic Unicode (e.g. 𝛼) to its plain form (α) so symbols in
# the equation match symbols in the legend regardless of styling.

_FUNC_NAMES = {
    "sin", "cos", "tan", "sec", "csc", "cot",
    "sinh", "cosh", "tanh", "log", "ln", "exp",
    "sqrt", "min", "max", "lim", "inf", "sup",
    "arg", "det", "mod", "gcd", "lcm",
}

# Word characters excluding digits and underscore — matches alpha runs in any
# script (Latin, Greek, math-italic …). We then drop any token that is purely
# Hangul or otherwise non-math.
_ALPHA_RUN_RE = re.compile(r"[^\W\d_]+", re.UNICODE)
_HANGUL_ONLY_RE = re.compile(r"^[가-힯]+$")


def normalize_symbol(s: str) -> str:
    """NFKD-normalize so math-italic chars (𝛼, 𝐴, 𝑎) collapse to plain (α, A, a)."""
    return unicodedata.normalize("NFKD", s)


def extract_equation_variables(omath_el) -> list[str]:
    """Identifier tokens (variable names) inside one <m:oMath> element,
    in first-occurrence order, deduplicated.

    Filters operators, numbers, Hangul labels, and well-known math function
    names (sin, log, …). The returned strings are kept in their original form
    (not normalized) for display purposes; matching against legend symbols
    should use ``normalize_symbol`` on both sides.
    """
    seen: set[str] = set()
    result: list[str] = []
    for el in omath_el.iter():
        local = el.tag.split("}")[-1] if "}" in el.tag else el.tag
        if local != "t" or not el.text:
            continue
        for tok in _ALPHA_RUN_RE.findall(el.text):
            if _HANGUL_ONLY_RE.match(tok):
                continue
            if normalize_symbol(tok).lower() in _FUNC_NAMES:
                continue
            if tok in seen:
                continue
            seen.add(tok)
            result.append(tok)
    return result


def equations_in_paragraph(para) -> list:
    """All top-level <m:oMath> elements inside a paragraph."""
    return _top_level_math_elements(para)


# ---------------------------------------------------------------------------
# Text post-processing
# ---------------------------------------------------------------------------

_UNICODE_NORMALIZE_MAP = str.maketrans({
    '‑': '-', '‐': '-', '‒': '-', '–': '-', '—': '-',
    ' ': ' ', ' ': ' ', ' ': ' ',
    '‘': "'", '’': "'", '“': '"', '”': '"',
    '…': '...',
})

_ABBREVS = {
    "fig", "figs", "e.g", "i.e", "no", "nos", "u.s", "vol", "approx",
    "cf", "vs", "et al", "sec", "art", "para", "ref", "dept",
}


def _normalize_unicode(text: str) -> str:
    return text.translate(_UNICODE_NORMALIZE_MAP)


def _break_sentences(text: str) -> str:
    protected = re.sub(
        r'\b(' + '|'.join(re.escape(a) for a in _ABBREVS) + r')\.',
        lambda m: m.group(1) + '\x00',
        text,
        flags=re.IGNORECASE,
    )
    broken = re.sub(r'\.\s+(?=[A-Z])', '.\n', protected)
    return broken.replace('\x00', '.')


# Indentation applied to every line that begins after a ';'-break in claim
# element lists and parameter legends. One tab character matches typical
# patent drafting style. The tab is emitted as a real ``<w:tab/>`` element
# by ``_build_text_run`` / ``write_run_with_breaks`` (Word's proper tab
# semantic) rather than as a literal whitespace character in ``<w:t>``.
_SEMICOLON_INDENT = "\t"


def _break_after_semicolons(text: str) -> str:
    return re.sub(r';\s+', f";\n{_SEMICOLON_INDENT}", text)


# ---------------------------------------------------------------------------
# 'respectively' expansion
# ---------------------------------------------------------------------------
# Defense-in-depth for the parameter-legend layout. The prompt already tells
# the LLM not to collapse parameters into 'A, B, C are X, Y, Z, respectively',
# but if it slips through we rewrite into per-parameter clauses joined by ';'.
#
# Conservative on purpose: we only fire when
#   - the symbol list is clearly math-like (short identifiers / Greek letters),
#   - each value is a simple noun phrase with no internal commas or semicolons,
#   - the symbol count and value count match,
#   - the word 'respectively' is either after the value list or before the verb.
# Anything fancier is left alone — better to skip a repair than to mangle text.

# Identifier shape: 1–10 chars, starts with a letter (ASCII or Greek), allows
# digits, subscripts, primes, underscores. Subscript characters U+2080..U+2089.
_SYM = r'[A-Za-zΑ-Ωα-ω][A-Za-z0-9Α-Ωα-ω₀-₉_\']{0,9}'
_VAL = r"[^,;.]+"

_VERB = (
    r'are|denote|denotes|represent|represents|stand\s+for|stands\s+for|'
    r'indicate|indicates|mean|means|refer\s+to|refers\s+to|'
    r'correspond\s+to|corresponds\s+to'
)

_LEGEND_VERB = (
    r'is|are|denote|denotes|represent|represents|stand\s+for|stands\s+for|'
    r'indicate|indicates|mean|means|refer\s+to|refers\s+to|'
    r'correspond\s+to|corresponds\s+to'
)

_RESPECTIVELY_AFTER_VALUES_RE = re.compile(
    rf'(?P<symbols>{_SYM}'
    rf'(?:\s*,\s*{_SYM}){{1,8}}'
    rf'(?:\s*,?\s*and\s+{_SYM})?)'
    rf'\s+(?P<verb>{_VERB})\s+'
    rf'(?P<values>{_VAL}'
    rf'(?:\s*,\s*{_VAL}){{1,8}}'
    rf'(?:\s*,?\s*and\s+{_VAL})?)'
    r'\s*,?\s*respectively',
    flags=re.IGNORECASE,
)

_RESPECTIVELY_BEFORE_VERB_RE = re.compile(
    rf'(?P<symbols>{_SYM}'
    rf'(?:\s*,\s*{_SYM}){{1,8}}'
    rf'(?:\s*,?\s*and\s+{_SYM})?)'
    r'\s+respectively\s+'
    rf'(?P<verb>{_VERB})\s+'
    rf'(?P<values>{_VAL}'
    rf'(?:\s*,\s*{_VAL}){{1,8}}'
    rf'(?:\s*,?\s*and\s+{_VAL})?)',
    flags=re.IGNORECASE,
)

_VERB_SINGULAR = {
    "is": "is",
    "are": "is",
    "denote": "denotes",
    "denotes": "denotes",
    "represent": "represents",
    "represents": "represents",
    "stand for": "stands for",
    "stands for": "stands for",
    "indicate": "indicates",
    "indicates": "indicates",
    "mean": "means",
    "means": "means",
    "refer to": "refers to",
    "refers to": "refers to",
    "correspond to": "corresponds to",
    "corresponds to": "corresponds to",
}

_MALFORMED_SEMICOLON_LEGEND_RE = re.compile(
    rf'(?P<prefix>\bwhere(?:in)?\s+)'
    rf'(?P<symbols>{_SYM}'
    rf'(?:\s*,\s*{_SYM}){{1,8}}'
    rf'(?:\s*,?\s*and\s+{_SYM})?)'
    rf'\s+(?P<verb>{_LEGEND_VERB})\s+'
    rf'(?P<first>{_VAL})'
    rf'(?P<rest>(?:\s*;\s*(?:where(?:in)?\s+)?'
    rf'(?:{_LEGEND_VERB})\s+{_VAL}){{1,8}})',
    flags=re.IGNORECASE,
)

_BARE_LEGEND_CLAUSE_RE = re.compile(
    rf'^\s*;\s*(?:where(?:in)?\s+)?(?P<verb>{_LEGEND_VERB})\s+(?P<value>{_VAL})',
    flags=re.IGNORECASE,
)


def _split_list(text: str) -> list[str]:
    parts = re.split(r'\s*,\s*and\s+|\s+and\s+|\s*,\s*', text.strip())
    return [p.strip() for p in parts if p.strip()]


def _expand_respectively(text: str) -> str:
    def repl(m: re.Match) -> str:
        symbols = _split_list(m.group("symbols"))
        values = _split_list(m.group("values"))
        if len(symbols) != len(values) or len(symbols) < 2:
            return m.group(0)
        verb = re.sub(r'\s+', ' ', m.group("verb").lower())
        singular = _VERB_SINGULAR.get(verb, verb)
        clauses = [f"{s} {singular} {v}" for s, v in zip(symbols, values)]
        return "; ".join(clauses)

    text = _RESPECTIVELY_AFTER_VALUES_RE.sub(repl, text)
    return _RESPECTIVELY_BEFORE_VERB_RE.sub(repl, text)


def _repair_malformed_semicolon_legend(text: str) -> str:
    """Repair 'where A, B, and C is X; is Y; is Z' style legends."""

    def repl(m: re.Match) -> str:
        symbols = _split_list(m.group("symbols"))
        values = [m.group("first").strip()]
        verbs = [m.group("verb")]
        for raw_clause in re.findall(r'\s*;\s*(?:where(?:in)?\s+)?(?:' + _LEGEND_VERB + r')\s+[^;.]+', m.group("rest"), flags=re.IGNORECASE):
            clause_match = _BARE_LEGEND_CLAUSE_RE.match(raw_clause)
            if clause_match is None:
                return m.group(0)
            verbs.append(clause_match.group("verb"))
            values.append(clause_match.group("value").strip())

        if len(symbols) != len(values) or len(symbols) < 2:
            return m.group(0)

        clauses: list[str] = []
        for i, (symbol, verb, value) in enumerate(zip(symbols, verbs, values)):
            normalized_verb = re.sub(r'\s+', ' ', verb.lower())
            singular = _VERB_SINGULAR.get(normalized_verb, normalized_verb)
            prefix = m.group("prefix") if i == 0 else ""
            clauses.append(f"{prefix}{symbol} {singular} {value}")
        return "; ".join(clauses)

    return _MALFORMED_SEMICOLON_LEGEND_RE.sub(repl, text)


_FIGS_REF_RE = re.compile(
    r'\b(?:figures|figs)\.?\s+(\d+[A-Za-z]?(?:\s*(?:,|and|to|-)\s*\d+[A-Za-z]?)+)',
    re.IGNORECASE,
)
_FIG_REF_RE = re.compile(r'\b(?:figure|fig)\.?\s+(\d+[A-Za-z]?)', re.IGNORECASE)
_BARE_FIG_RE = re.compile(r'\bFIG\s+(\d+[A-Za-z]?)\b')


def _normalize_figure_refs(text: str) -> str:
    """Normalize figure references to USPTO-style FIG./FIGS."""

    def plural_repl(m: re.Match) -> str:
        refs = re.sub(
            r'\d+[A-Za-z]?',
            lambda ref: ref.group(0).upper(),
            m.group(1),
        )
        return f"FIGS. {refs}"

    def singular_repl(m: re.Match) -> str:
        return f"FIG. {m.group(1).upper()}"

    text = _FIGS_REF_RE.sub(plural_repl, text)
    text = _FIG_REF_RE.sub(singular_repl, text)
    return _BARE_FIG_RE.sub(singular_repl, text)


def postprocess(text: str) -> str:
    text = _normalize_unicode(text)
    text = _normalize_figure_refs(text)
    text = _expand_respectively(text)
    text = _repair_malformed_semicolon_legend(text)
    text = _break_sentences(text)
    text = _break_after_semicolons(text)
    return text
